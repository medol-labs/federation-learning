#!/usr/bin/env python3
import argparse
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

BODY_FONT = "PingFang SC"
MONO_FONT = "Consolas"


def parse_args():
    parser = argparse.ArgumentParser(description="Convert project Markdown to DOCX.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--reference", type=Path)
    parser.add_argument("--title", default="")
    return parser.parse_args()


def main():
    args = parse_args()
    markdown = args.input.read_text(encoding="utf-8")
    document = Document()
    use_landscape = "数据库设计说明" in args.title or "数据库设计说明" in args.input.name
    configure_document(document, args.reference, use_landscape)
    write_markdown(document, markdown, args.title)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.out)


def configure_document(document, reference, use_landscape=False):
    section = document.sections[0]
    if reference and reference.exists():
        ref = Document(reference)
        ref_section = ref.sections[0]
        section.page_width = ref_section.page_width
        section.page_height = ref_section.page_height
        section.top_margin = ref_section.top_margin
        section.bottom_margin = ref_section.bottom_margin
        section.left_margin = ref_section.left_margin
        section.right_margin = ref_section.right_margin
        section.header_distance = ref_section.header_distance
        section.footer_distance = ref_section.footer_distance
    else:
        section.page_width = Cm(21.59)
        section.page_height = Cm(27.94)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    if use_landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        if section.page_width < section.page_height:
            section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Cm(1.4)
        section.right_margin = Cm(1.4)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    set_style_font(title, BODY_FONT)
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    remove_paragraph_borders(title.element)

    for index in range(1, 7):
        style = styles[f"Heading {index}"]
        set_style_font(style, BODY_FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.font.size = Pt(max(16 - index, 10.5))
        style.paragraph_format.space_before = Pt(10 if index <= 2 else 6)
        style.paragraph_format.space_after = Pt(6)


def write_markdown(document, markdown, fallback_title):
    lines = strip_html_comments(markdown).splitlines()
    index = 0
    wrote_title = False
    in_code = False
    code_language = ""
    code_lines = []

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                code = "\n".join(code_lines)
                if code_language == "mermaid":
                    add_mermaid_diagram(document, code)
                else:
                    add_code_block(document, code)
                code_lines = []
                in_code = False
                code_language = ""
            else:
                in_code = True
                code_language = stripped.removeprefix("```").strip().lower()
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1 and not wrote_title:
                paragraph = document.add_paragraph(text, style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                wrote_title = True
            else:
                document.add_paragraph(text, style=f"Heading {min(level, 6)}")
            index += 1
            continue

        if is_table_start(lines, index):
            rows, index = collect_table(lines, index)
            add_table(document, rows)
            continue

        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        ordered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if bullet or ordered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline_runs(paragraph, (bullet or ordered).group(1).strip())
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            lookahead = lines[index].strip()
            if not lookahead or lookahead.startswith("#") or lookahead.startswith("```") or is_table_start(lines, index):
                break
            if re.match(r"^\s*([-*]|\d+\.)\s+", lines[index]):
                break
            paragraph_lines.append(lookahead)
            index += 1
        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, " ".join(paragraph_lines))

    if not wrote_title and fallback_title:
        first = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
        first.text = fallback_title
        first.style = "Title"


def set_style_font(style, font_name):
    style.font.name = font_name
    if style._element.rPr is None:
        style._element.get_or_add_rPr()
    r_fonts = style._element.rPr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        style._element.rPr.append(r_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), font_name)


def set_run_font(run, font_name=BODY_FONT):
    run.font.name = font_name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    r_fonts = run._element.rPr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        run._element.rPr.append(r_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), font_name)


def remove_paragraph_borders(element):
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def strip_html_comments(markdown):
    return re.sub(r"<!--[\s\S]*?-->", "", markdown)


def clean_inline(value):
    return re.sub(r"`([^`]+)`", r"\1", value).replace("**", "").strip()


def add_inline_runs(paragraph, text):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            set_run_font(paragraph.add_run(text[position:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO_FONT)
        position = match.end()
    if position < len(text):
        set_run_font(paragraph.add_run(text[position:]))


def add_code_block(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(code)
    set_run_font(run, MONO_FONT)
    run.font.size = Pt(9)


def add_mermaid_diagram(document, code):
    image = render_mermaid_flowchart(code)
    if image is None:
        add_code_block(document, code)
        return
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp:
        image.save(temp.name)
        temp_path = temp.name
    try:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(temp_path, width=Cm(16.5))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        caption_run = caption.add_run("图 2-1 联邦学习子系统部署拓扑")
        set_run_font(caption_run)
        caption_run.font.size = Pt(9)
        caption_run.italic = True
    finally:
        Path(temp_path).unlink(missing_ok=True)


def render_mermaid_flowchart(code):
    nodes = {}
    edges = []
    node_pattern = re.compile(r'^\s*([A-Za-z][A-Za-z0-9_]*)\["(.+)"\]\s*$')
    edge_pattern = re.compile(r'^\s*([A-Za-z][A-Za-z0-9_]*)\s*-->\s*([A-Za-z][A-Za-z0-9_]*)\s*$')
    for line in code.splitlines():
        node_match = node_pattern.match(line)
        if node_match:
            nodes[node_match.group(1)] = node_match.group(2)
            continue
        edge_match = edge_pattern.match(line)
        if edge_match:
            edges.append((edge_match.group(1), edge_match.group(2)))
    if not nodes or not edges:
        return None

    width, height = 1800, 1050
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = load_diagram_font(32)
    small_font = load_diagram_font(26)

    layer_ids = [
        ["用户"],
        ["Gateway"],
        [node_id for node_id in nodes if node_id.startswith("Web")],
        [node_id for node_id in nodes if node_id.startswith("Api")],
        [node_id for node_id in nodes if node_id in {"Db", "Storage", "RuntimeEngine", "External"}],
    ]
    placed = {node_id for layer in layer_ids for node_id in layer}
    remaining = [node_id for node_id in nodes if node_id not in placed]
    if remaining:
        layer_ids.append(remaining)

    positions = {}
    box_width, box_height = 310, 92
    layer_y = [70, 210, 365, 560, 795, 930]
    for layer_index, ids in enumerate(layer_ids):
        ids = [node_id for node_id in ids if node_id in nodes]
        if not ids:
            continue
        y = layer_y[min(layer_index, len(layer_y) - 1)]
        gap = (width - 2 * 120 - len(ids) * box_width) / max(len(ids) - 1, 1)
        if len(ids) == 1:
            xs = [(width - box_width) / 2]
        else:
            xs = [120 + index * (box_width + gap) for index in range(len(ids))]
        for node_id, x in zip(ids, xs):
            positions[node_id] = (int(x), int(y), int(x + box_width), int(y + box_height))

    for source, target in edges:
        if source in positions and target in positions:
            draw_arrow(draw, positions[source], positions[target])

    for node_id, rect in positions.items():
        fill = "#EAF4FB"
        outline = "#8EB8D9"
        if node_id == "Gateway":
            fill = "#E8F1FF"
            outline = "#5E8AC7"
        elif node_id.startswith("Api"):
            fill = "#EAF7EE"
            outline = "#75B88A"
        elif node_id in {"Db", "Storage"}:
            fill = "#FFF5E6"
            outline = "#D1A35B"
        elif node_id in {"RuntimeEngine", "External"}:
            fill = "#F3F0FF"
            outline = "#9A85C7"
        draw.rounded_rectangle(rect, radius=18, fill=fill, outline=outline, width=3)
        label = normalize_diagram_label(nodes[node_id])
        draw_wrapped_centered_text(draw, label, rect, font if len(label) <= 16 else small_font)

    title_font = load_diagram_font(34)
    draw.text((width / 2, 24), "联邦学习子系统部署拓扑", font=title_font, fill="#111111", anchor="mm")
    return image


def load_diagram_font(size):
    font_candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for path in font_candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size)
            except OSError:
                continue
    return ImageFont.load_default()


def normalize_diagram_label(label):
    replacements = {
        "联合 Learning Console 前端（平台控制台）": "平台控制台前端",
        "联合 Learning Participant Console 前端（参与方控制台）": "参与方控制台前端",
        "联邦学习支撑服务 后端（支撑服务）": "支撑服务后端",
        "联邦学习平台 后端（平台控制面）": "平台控制面后端",
        "联邦学习运行时代理 后端（参与方 Runtime Agent）": "参与方运行时代理",
        "运行时 引擎": "运行时引擎",
    }
    return replacements.get(label, label)


def draw_arrow(draw, source_rect, target_rect):
    sx = (source_rect[0] + source_rect[2]) / 2
    sy = source_rect[3]
    tx = (target_rect[0] + target_rect[2]) / 2
    ty = target_rect[1]
    if abs(tx - sx) < 40 and ty >= sy:
        points = [(sx, sy), (tx, ty)]
    else:
        mid_y = (sy + ty) / 2
        points = [(sx, sy), (sx, mid_y), (tx, mid_y), (tx, ty)]
    draw.line(points, fill="#667085", width=4, joint="curve")
    draw.polygon([(tx, ty), (tx - 10, ty - 18), (tx + 10, ty - 18)], fill="#667085")


def draw_wrapped_centered_text(draw, text, rect, font):
    max_width = rect[2] - rect[0] - 28
    lines = wrap_text(draw, text, font, max_width)
    line_heights = [draw.textbbox((0, 0), line, font=font)[3] - draw.textbbox((0, 0), line, font=font)[1] for line in lines]
    total_height = sum(line_heights) + max(0, len(lines) - 1) * 6
    y = rect[1] + (rect[3] - rect[1] - total_height) / 2
    for line, line_height in zip(lines, line_heights):
        draw.text(((rect[0] + rect[2]) / 2, y), line, font=font, fill="#111111", anchor="ma")
        y += line_height + 6


def wrap_text(draw, text, font, max_width):
    lines = []
    current = ""
    for char in text:
        candidate = current + char
        if current and draw.textlength(candidate, font=font) > max_width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines[:3]


def is_table_start(lines, index):
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    separator = lines[index + 1].strip()
    return current.startswith("|") and current.endswith("|") and is_table_separator(separator)


def collect_table(lines, index):
    rows = []
    while index < len(lines):
        line = lines[index].strip()
        if not line.startswith("|") or not line.endswith("|"):
            break
        if not is_table_separator(line):
            rows.append([clean_inline(cell.strip()) for cell in line.strip("|").split("|")])
        index += 1
    return rows, index


def is_table_separator(line):
    if not line.startswith("|") or not line.endswith("|"):
        return False
    cells = [cell.strip() for cell in line.strip("|").split("|")]
    return bool(cells) and all(re.match(r"^:?-{3,}:?$", cell) for cell in cells)


def add_table(document, rows):
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    is_wide = column_count >= 4
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = not is_wide
    widths = table_column_widths(column_count) if is_wide else []
    for row_index, row in enumerate(rows):
        row_cells = table.rows[row_index].cells
        if row_index == 0:
            repeat_table_header(table.rows[row_index])
        for column_index in range(column_count):
            cell = row_cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell, top=80 if is_wide else 90, start=90 if is_wide else 120, bottom=80 if is_wide else 90, end=90 if is_wide else 120)
            if widths:
                set_cell_width(cell, widths[column_index])
            if row_index == 0:
                set_cell_fill(cell, "D9EAF7")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1 if is_wide else 1.15
            text = row[column_index] if column_index < len(row) else ""
            add_inline_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(8 if is_wide else 9.5)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
            if len(text) < 20:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()


def table_column_widths(column_count):
    if column_count == 4:
        return [Cm(4.6), Cm(3.8), Cm(3.8), Cm(13.0)]
    if column_count == 5:
        return [Cm(4.2), Cm(3.6), Cm(3.2), Cm(2.6), Cm(12.0)]
    if column_count == 6:
        return [Cm(3.4), Cm(3.2), Cm(2.1), Cm(2.2), Cm(2.1), Cm(11.8)]
    return [Cm(24.8 / column_count) for _ in range(column_count)]


def set_cell_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width.twips))
    tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def set_cell_margins(cell, top, start, bottom, end):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


if __name__ == "__main__":
    main()
