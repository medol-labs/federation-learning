import unittest

from docx import Document

from markdown_to_docx import write_markdown


class HeadingNumberingTest(unittest.TestCase):
    def test_merge_chapter_numbers_table_then_business_groups(self):
        markdown = (
            "# 3.3.1 联邦学习平台模块配置项功能需求\n\n"
            "## 表 3 联邦学习模块的软件功能划分\n\n"
            "## 组织管理\n\n"
            "## 联邦管理\n"
        )
        document = Document()
        write_markdown(document, markdown, "", "3.3.1")
        headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
        self.assertEqual(headings, [
            "3.3.1.1 表 3 联邦学习模块的软件功能划分",
            "3.3.1.2 组织管理",
            "3.3.1.3 联邦管理",
        ])

    def test_numbers_subheadings_only_when_prefix_is_requested(self):
        markdown = "# 3.3.1 模块需求\n\n## 组织管理\n\n### 注册组织\n\n## 联邦管理\n"
        numbered = Document()
        write_markdown(numbered, markdown, "", "3.3.1")
        headings = [paragraph.text for paragraph in numbered.paragraphs if paragraph.style.name.startswith("Heading")]
        self.assertEqual(headings, [
            "3.3.1.1 组织管理",
            "3.3.1.1.1 注册组织",
            "3.3.1.2 联邦管理",
        ])

        plain = Document()
        write_markdown(plain, markdown, "")
        self.assertEqual(
            [paragraph.text for paragraph in plain.paragraphs if paragraph.style.name.startswith("Heading")],
            ["组织管理", "注册组织", "联邦管理"],
        )


if __name__ == "__main__":
    unittest.main()
