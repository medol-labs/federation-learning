"""Checks for controlled corrections to the joint requirements overview."""

import importlib.util
import json
import sys
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'scripts/merge-438c-project-documents.py'
sys.path.insert(0, str(ROOT / 'scripts'))
spec = importlib.util.spec_from_file_location('merge_438c_project_documents', SCRIPT)
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)


class RequirementsOverviewCorrectionsTest(unittest.TestCase):
    def setUp(self):
        source = ROOT / 'docs/438c/merge/requirements-overview-corrections.json'
        self.items = json.loads(source.read_text(encoding='utf-8'))['items']
        self.document = Document()
        for item in self.items:
            self.document.add_paragraph(item['match'])

    def test_replaces_only_the_reviewed_paragraphs(self):
        merge.apply_requirements_overview_corrections(self.document, check=True)
        self.assertEqual([p.text for p in self.document.paragraphs], [item['match'] for item in self.items])
        merge.apply_requirements_overview_corrections(self.document, check=False)
        self.assertEqual([p.text for p in self.document.paragraphs], [item['replacement'] for item in self.items])

    def test_stops_if_joint_source_changes(self):
        self.document.paragraphs[0].text = 'Changed source paragraph'
        with self.assertRaisesRegex(ValueError, 'FL-RQ-OVERVIEW-01'):
            merge.apply_requirements_overview_corrections(self.document, check=True)


if __name__ == '__main__':
    unittest.main()
