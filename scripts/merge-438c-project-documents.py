#!/usr/bin/env python3
"""Produce review copies of the other joint 438C documents with FL sections inserted."""

import argparse
import hashlib
import importlib.util
import json
import re
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from docx.shared import Twips

from database_design_inventory import (
    load_catalog, render_logical, render_physical, traceability_rows,
)


ROOT = Path(__file__).resolve().parents[1]
TARGET_DIR = ROOT.parent / 'tmp/merge/项目文档v0.2'
DATABASE_TARGET = ROOT.parent / 'tmp/merge/03“感一通一算”一体化的跨域健康管理联邦孪生平台-数据库设计说明v1.0.docx'
SOURCE_DIR = ROOT / 'docs/438c/merge'
OUTPUT_DIR = ROOT / 'docs/generated/438c/merge/project-documents'
SOFTWARE_SCRIPT = ROOT / 'scripts/merge-438c-software-design.py'
spec = importlib.util.spec_from_file_location('software_design_merge', SOFTWARE_SCRIPT)
software_merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(software_merge)
formatter_spec = importlib.util.spec_from_file_location('requirement_formatter', ROOT / 'scripts/format-438c-331.py')
requirement_formatter = importlib.util.module_from_spec(formatter_spec)
formatter_spec.loader.exec_module(requirement_formatter)
insert_section = software_merge.insert_section
from markdown_to_docx import format_requirements_tables

MARKER = re.compile(r'^<!-- INSERT ([a-z-]+) -->$', re.MULTILINE)


def sections(name, expected):
    text = (SOURCE_DIR / name).read_text(encoding='utf-8')
    matches = list(MARKER.finditer(text))
    if len(matches) != len(expected) or set(match.group(1) for match in matches) != set(expected):
        raise ValueError(f'{name}: missing or duplicate insertion markers')
    return {
        match.group(1): text[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(text)].strip()
        for index, match in enumerate(matches)
    }


def heading(document, text, style, parent=None):
    current_parent = None
    found = []
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else None
        if style_name == 'Heading 2':
            current_parent = paragraph.text.strip()
        if paragraph.text.strip() == text and style_name == style:
            if parent is None or current_parent == parent:
                found.append(paragraph)
    if len(found) != 1:
        raise ValueError(f'Expected one {style} heading {text!r} under {parent!r}; found {len(found)}')
    return found[0]


def convert_requirements(source, out_dir):
    command = [
        'soffice', f'-env:UserInstallation=file://{out_dir / "lo-profile"}',
        '--headless', '--convert-to', 'docx', '--outdir', str(out_dir), str(source),
    ]
    result = subprocess.run(command, capture_output=True, text=True, check=False)
    converted = out_dir / (source.stem + '.docx')
    if result.returncode or not converted.exists():
        raise RuntimeError(f'Could not convert the joint requirements document: {result.stderr}')
    return converted


def requirement_parts():
    draft = (SOURCE_DIR / '3.3.1-federation-learning.md').read_text(encoding='utf-8')
    register = json.loads((SOURCE_DIR / '3.3.1-function-register.json').read_text(encoding='utf-8'))
    active_count = sum(item.get('status') != 'retired' for item in register['functions'])
    source = requirement_formatter.format_document(draft)
    lines = source.splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith('| 序号 |'))
    model_lines = [line for line in lines[start:] if line.startswith('|')]
    from markdown_to_docx import collect_table
    rows, end = collect_table(lines, start)
    if len(rows) != active_count + 1 or len(rows[0]) != 5:
        raise ValueError(f'Expected header plus {active_count} requirements, got {len(rows)} rows')
    if len(model_lines) < len(rows):
        raise ValueError('Requirement table is incomplete')
    intro = '\n'.join(lines[2:next(i for i, line in enumerate(lines) if line.startswith('## 表 3'))]).strip()
    intro = re.sub(r'^- (.+)$', r'· \1\n', intro, flags=re.MULTILINE)
    details = '\n'.join(lines[end:]).strip()
    details = re.sub(r'^## ', '#### ', details, flags=re.MULTILINE)
    details = re.sub(r'^- (.+)$', r'· \1\n', details, flags=re.MULTILINE)
    if not details.startswith('#### 组织管理'):
        raise ValueError('Requirement detail section has an unexpected first heading')
    return intro, rows, details


def requirement_placeholder(document):
    anchor = heading(document, '联邦学习模块配置项功能需求/LBLS-GN-LBXX', 'Heading 3')
    objects = []
    table = None
    for element in anchor._p.itersiblings():
        if element.tag == qn('w:p'):
            paragraph = Paragraph(element, document._body)
            if paragraph.style is not None and paragraph.style.name == 'Heading 3':
                break
            objects.append(paragraph)
        elif element.tag == qn('w:tbl'):
            if table is not None:
                raise ValueError('More than one requirement placeholder table')
            table = next(item for item in document.tables if item._tbl is element)
    if table is None or len(table.rows) != 2 or len(table.columns) != 5:
        raise ValueError('Requirement placeholder table has changed')
    expected_placeholders = {'所有功能标识', '挂*****/LBLS-GN-LBXX-****'}
    found = {paragraph.text.strip() for paragraph in objects if paragraph.style is not None and paragraph.style.name == 'Heading 4'}
    if found != expected_placeholders:
        raise ValueError(f'Requirement placeholder headings have changed: {found}')
    return anchor, objects, table


def merge_requirements(target, out, check):
    document = Document(target)
    apply_requirements_overview_corrections(document, check)
    anchor, objects, table = requirement_placeholder(document)
    intro, rows, details = requirement_parts()
    if check:
        return
    existing_tables = {item._tbl for item in document.tables}
    for paragraph in objects:
        if paragraph.style is not None and paragraph.style.name == 'Heading 4':
            paragraph._p.getparent().remove(paragraph._p)
    insert_section(document, anchor, intro)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for fields in rows[1:]:
        row = table.add_row()
        for cell, value in zip(row.cells, fields):
            cell.text = value
    marker = document.add_paragraph()
    table._tbl.addnext(marker._p)
    insert_section(document, marker, details)
    marker._p.getparent().remove(marker._p)
    added_tables = [item for item in document.tables if item._tbl not in existing_tables]
    format_requirements_tables(document, [table, *added_tables])
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def apply_requirements_overview_corrections(document, check):
    source = json.loads((SOURCE_DIR / 'requirements-overview-corrections.json').read_text(encoding='utf-8'))
    if source.get('schemaVersion') != 1:
        raise ValueError('Unsupported requirements overview corrections schema')
    for item in source['items']:
        matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == item['match']]
        if len(matches) != 1:
            raise ValueError(f'Requirements overview source changed for {item["id"]}: found {len(matches)} matches')
        if not check:
            paragraph = matches[0]
            paragraph.runs[0].text = item['replacement']
            for run in paragraph.runs[1:]:
                run.text = ''


def merge_database(target, out, check):
    baseline = json.loads((SOURCE_DIR / 'database-design-merge-base.json').read_text(encoding='utf-8'))
    if baseline.get('schemaVersion') != 1 or target.resolve() != DATABASE_TARGET.resolve():
        raise ValueError('Database design must use the reviewed joint source document')
    actual_hash = hashlib.sha256(target.read_bytes()).hexdigest()
    if actual_hash != baseline['sha256']:
        raise ValueError('Joint database design source changed; review and update database-design-merge-base.json')
    content = sections('database-design-federation-learning.md', {
        'conceptual', 'logical', 'physical', 'software-units',
    })
    catalog = load_catalog()
    requirement_register = json.loads((SOURCE_DIR / '3.3.1-function-register.json').read_text(encoding='utf-8'))
    logical_inventory, logical_details = render_logical(catalog, requirement_register)
    entity_inventory, entity_details = render_physical(catalog)
    content['logical'] = content['logical'].replace('{{READ_MODEL_INVENTORY}}', logical_inventory)
    content['logical'] = content['logical'].replace('{{READ_MODEL_DETAILS}}', logical_details)
    content['physical'] = content['physical'].replace('{{ENTITY_INVENTORY}}', entity_inventory)
    content['physical'] = content['physical'].replace('{{ENTITY_DETAILS}}', entity_details)
    trace_rows = traceability_rows(catalog, requirement_register)
    for key, text in content.items():
        if re.search(r'\{\{(?:READ_MODEL|ENTITY|DATABASE|TRACEABILITY)_[A-Z_]+\}\}', text):
            raise ValueError(f'Unresolved database design placeholder: {key}')
    document = Document(target)
    anchors = {
        key: heading(document, '联邦学习模块', 'Heading 3', parent)
        for key, parent in {
            'conceptual': '数据库概念设计',
            'logical': '数据库逻辑设计',
            'physical': '数据库物理设计',
        }.items()
    }
    software_units = heading(document, '用于数据库访问或操纵的软件单元的详细设计', 'Heading 1')
    software_table = following_table(document, software_units)
    if software_table.rows[0].cells[0].text != '软件单元名称':
        raise ValueError('Database software-unit table format changed')
    software_intro = next(
        (paragraph for paragraph in document.paragraphs
         if paragraph.text.strip() == '跨域健康管理联邦孪生平台的软件单元设计见表 117。'),
        None,
    )
    if software_intro is None:
        raise ValueError('Joint software-unit introduction changed')
    corrections = json.loads((SOURCE_DIR / 'database-design-corrections.json').read_text(encoding='utf-8'))
    if corrections.get('schemaVersion') != 1:
        raise ValueError('Unsupported database design corrections schema')
    for correction in corrections['items']:
        matches = [p for p in document.paragraphs if correction['match'] in p.text]
        if len(matches) != 1:
            raise ValueError(f'{correction["id"]}: expected one matching paragraph, found {len(matches)}')
    figures = {
        'FL-DB-FIG-GOVERNANCE': (
            {'widthCm': 12.5, 'caption': '联邦学习组织与数据治理概念关系图'},
            SOURCE_DIR / 'figures/conceptual-data-governance.png',
        ),
        'FL-DB-FIG-TRAINING': (
            {'widthCm': 9.5, 'caption': '联邦学习训练与模型更新概念关系图'},
            SOURCE_DIR / 'figures/conceptual-data-training.png',
        ),
    }
    if any(not path.is_file() for _, path in figures.values()):
        raise ValueError('Database conceptual figure is missing')
    forward = heading(document, '正向追溯', 'Heading 2')
    reverse = heading(document, '逆向追溯', 'Heading 2')
    forward_table = following_table(document, forward)
    reverse_table = following_table(document, reverse)
    if len(forward_table.columns) != 4 or len(reverse_table.columns) != 4:
        raise ValueError('Database traceability table format changed')
    validate_database_traceability_table(forward_table, True)
    validate_database_traceability_table(reverse_table, False)
    if check:
        return
    for correction in corrections['items']:
        paragraph = next(p for p in document.paragraphs if correction['match'] in p.text)
        paragraph.text = paragraph.text.replace(correction['match'], correction['replacement'])
    software_intro.text = '跨域健康管理联邦孪生平台的软件单元设计'
    software_intro.style = document.styles['Heading 2']
    old_tables = {table._tbl for table in document.tables}
    for key, anchor in anchors.items():
        insert_section(document, anchor, content[key], figures if key == 'conceptual' else None)
    marker = document.add_paragraph()
    software_table._tbl.addnext(marker._p)
    insert_section(document, marker, content['software-units'])
    marker._p.getparent().remove(marker._p)
    training_figure_heading = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip() == '训练链路概念关系图'
    )
    training_figure_heading.paragraph_format.keep_with_next = True
    merge_database_traceability_table(forward_table, trace_rows, True)
    merge_database_traceability_table(reverse_table, trace_rows, False)
    empty_captions = []
    within_trace = False
    for paragraph in document.paragraphs:
        if paragraph._p is forward._p:
            within_trace = True
        elif paragraph._p is reverse._p:
            break
        elif within_trace and paragraph.style.name == 'RS-表-题注' and not paragraph.text.strip():
            empty_captions.append(paragraph)
    if len(empty_captions) != 1:
        raise ValueError('Expected one empty caption after the joint forward trace table')
    empty_captions[0]._p.getparent().remove(empty_captions[0]._p)
    format_database_tables(document, [t for t in document.tables if t._tbl not in old_tables])
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def validate_database_traceability_table(table, forward):
    if len(table.columns) != 4 or len(table.rows) < 12 or table.rows[0].cells[0].text != '序号':
        raise ValueError('Database traceability table format changed')
    requirement_column = 1 if forward else 2
    if any('LBLS-GN-LBXX-' not in row.cells[requirement_column].text for row in table.rows[1:11]):
        raise ValueError('Expected ten contiguous federation-learning trace rows in joint source')
    if 'LBLS-GN-LBXX-' in table.rows[11].cells[requirement_column].text:
        raise ValueError('Federation-learning trace rows extend beyond the expected block')


def merge_database_traceability_table(table, trace_rows, forward):
    original_rows = table.rows[1:11]
    insertion_point = table.rows[11]._tr
    template_run = original_rows[0].cells[1 if forward else 2].paragraphs[0].runs[0]
    for requirement_id, title, logical_id, physical_id, _table_name in trace_rows:
        design_ref = f'{logical_id}；{physical_id}'
        requirement_ref = f'{title}/{requirement_id}'
        cells = ('', requirement_ref, design_ref, '') if forward else ('', design_ref, requirement_ref, '')
        new_row = _Row(deepcopy(original_rows[0]._tr), table)
        for cell, value in zip(new_row.cells, cells):
            set_table_cell_text(cell, value, template_run)
        insertion_point.addprevious(new_row._tr)
    for row in original_rows:
        row._tr.getparent().remove(row._tr)
    for index, row in enumerate(table.rows[1:], 1):
        set_table_cell_text(row.cells[0], str(index), template_run)
    table.autofit = False
    widths = (650, 3300, 4500, 616)
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)


def set_table_cell_text(cell, value, template_run):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        run = paragraph.add_run(value)
        if template_run._r.rPr is not None:
            run._r.insert(0, deepcopy(template_run._r.rPr))


def following_table(document, paragraph):
    for element in paragraph._p.itersiblings():
        if element.tag == qn('w:tbl'):
            return Table(element, document._body)
        if element.tag == qn('w:p') and Paragraph(element, document._body).style.name.startswith('Heading'):
            break
    raise ValueError(f'Missing table after {paragraph.text}')


def format_database_tables(document, tables):
    software_merge.format_inserted_tables(document, tables)
    column_widths = {
        ('设计标识', '部署单元', '查询视图', '逻辑标识', '对应需求'): [1600, 950, 1500, 2550, 2400],
        ('设计标识', '部署单元', 'JPA 查询表', '主键属性', '对应实体'): [1600, 950, 2350, 2000, 2100],
        ('数据元素', '逻辑类型', '数量', '必需', '标识'): [3300, 2200, 1000, 1000, 1500],
        ('JPA 属性', '映射类型', '主键', '长文本', '说明'): [3000, 2500, 950, 950, 1600],
        ('序号', '需求功能', '数据库设计标识与查询表'): [700, 3200, 5100],
        ('序号', '数据库设计标识与查询表', '对应需求功能'): [700, 5100, 3200],
    }
    for table in tables:
        widths = column_widths.get(tuple(cell.text.strip() for cell in table.rows[0].cells))
        if widths:
            for column, width in zip(table.columns, widths):
                column.width = Twips(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Twips(width)
        if table.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is None:
            header = OxmlElement('w:tblHeader')
            header.set(qn('w:val'), 'true')
            table.rows[0]._tr.get_or_add_trPr().append(header)


def merge_user_manual(target, out, check):
    content = sections('user-manual-federation-learning.md', {'procedure', 'scenario'})
    document = Document(target)
    procedure = heading(document, '联邦学习模块', 'Heading 3')
    scenario = heading(document, '多中心联邦建模科研场景', 'Heading 2')
    next_section = heading(document, '医疗数字孪生辅助临床诊疗场景', 'Heading 2')
    if check:
        return
    insert_section(document, procedure, content['procedure'])
    element = scenario._p.getnext()
    while element is not None and element is not next_section._p:
        following = element.getnext()
        element.getparent().remove(element)
        element = following
    insert_section(document, scenario, content['scenario'])
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def merge_installation(target, out, check):
    content = sections('installation-federation-learning.md', {'installation'})
    document = Document(target)
    heading(document, '系统部署安装', 'Heading 1')
    if check:
        return
    marker = document.add_paragraph()
    insert_section(document, marker, content['installation'])
    marker._p.getparent().remove(marker._p)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--target-dir', type=Path, default=TARGET_DIR)
    parser.add_argument('--out-dir', type=Path, default=OUTPUT_DIR)
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--only', choices=('01', '03', '04', '05'), help='Generate or check one document')
    args = parser.parse_args()
    codes = (args.only,) if args.only else ('01', '03', '04', '05')
    targets = {code: DATABASE_TARGET if code == '03' else next(args.target_dir.glob(f'{code}*')) for code in codes}
    if any(target.resolve() == (args.out_dir / f'{code}-federation-learning-merged.docx').resolve() for code, target in targets.items()):
        raise ValueError('The output must not overwrite a joint source document')
    with tempfile.TemporaryDirectory() as directory:
        requirements_docx = convert_requirements(targets['01'], Path(directory)) if '01' in targets else None
        jobs = [
            ('01', merge_requirements, requirements_docx),
            ('03', merge_database, targets.get('03')),
            ('04', merge_user_manual, targets.get('04')),
            ('05', merge_installation, targets.get('05')),
        ]
        for code, merge, target in jobs:
            if args.only is not None and code != args.only:
                continue
            output = args.out_dir / f'{code}-federation-learning-merged.docx'
            merge(target, output, args.check)
            print(f'{code}: {"checked" if args.check else output}')


if __name__ == '__main__':
    main()
