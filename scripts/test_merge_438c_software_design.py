"""Focused checks for the controlled software-design merge baseline."""

import importlib.util
import json
import re
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SCRIPT = Path(__file__).with_name('merge-438c-software-design.py')
spec = importlib.util.spec_from_file_location('merge_438c_software_design', SCRIPT)
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)


class SoftwareDesignTrackingTest(unittest.TestCase):
    def test_other_sections_render_from_stable_items(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        rendered = merge.render_structured_sections(source)
        self.assertEqual(set(rendered), set(merge.ANCHORS) - {'decisions'})
        self.assertIn('{{READ_MODEL_INVENTORY}}', rendered['db-logical'])
        self.assertIn('| 部件 | 主要职责 | 持久化或外部依赖 |', rendered['architecture'])
        self.assertIn('### 联邦学习模块性能设计', rendered['performance'])
        self.assertEqual(rendered['interfaces'].count('{{INTERFACE_INVENTORY}}'), 1)
        self.assertIn('当前各后端模块配置了 Hibernate ddl-auto=update', rendered['db-physical'])
        self.assertIn('当前平台聚合适配器会在事件处理链中轮询等待引擎结果', rendered['execution'])
        self.assertIn('当前引擎服务端不校验', rendered['architecture'])
        self.assertIn('各软件单元之间的通信与数据访问边界', rendered['architecture'])
        self.assertNotIn('### 安全边界与权限决策', rendered['architecture'])
        self.assertIn('FL-U-AGG 模型聚合', rendered['detailed'])
        self.assertIn('FL-U-EXEC 本地轮次执行', rendered['detailed'])
        self.assertIn('FL-U-C-ORG 组织管理 / FL-SD-DET-26-R01', rendered['detailed'])
        self.assertIn('FL-U-SEL 参与者选择 / FL-SD-DET-05-R01', rendered['detailed'])
        self.assertIn('注册组织 / FL-SD-OP-F0001', rendered['detailed'])
        self.assertIn('字典值翻译目录 / FL-SD-OP-F0171', rendered['detailed'])
        self.assertIn('服务账户 API 令牌目录 / FL-SD-OP-F0192', rendered['detailed'])
        self.assertIn('FL-U-C-IAM 身份访问管理 / FL-SD-DET-26-R15', rendered['detailed'])
        self.assertIn('当前本地全局模型评估适配器返回固定精度与公平性数值', rendered['detailed'])
        self.assertIn('LBLS-GN-LBXX-F0131', rendered['detailed'])
        self.assertIn('{{FIGURE:FL-SD-DBC-04}}', rendered['db-conceptual'])
        self.assertIn('{{FIGURE:FL-SD-DBC-05}}', rendered['db-conceptual'])
        self.assertIn('组织与联邦均可关联多条成员资格', rendered['db-conceptual'])

    def test_detailed_design_requirement_references_are_current(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        signatures = merge.requirement_signatures(source)
        self.assertEqual(len(signatures), 19)
        self.assertIn('LBLS-GN-LBXX-F0066', signatures)
        changed = deepcopy(source)
        table = next(block for block in changed['sections']['detailed'] if block['id'] == 'FL-SD-DET-23')
        table['rows'][0]['cells'][1] = '错误的功能名称'
        with self.assertRaisesRegex(ValueError, 'requirement trace is invalid'):
            merge.requirement_signatures(changed)

    def test_operation_design_covers_commands_queries_and_automation(self):
        groups = merge.load_operation_design()
        rows = {row['requirementId']: row for group in groups for row in group['rows']}
        self.assertEqual(len(rows), 192)
        self.assertIn('组织名称必须唯一', rows['LBLS-GN-LBXX-F0001']['cells'][2])
        self.assertIn('查询视图按条件读取', rows['LBLS-GN-LBXX-F0005']['cells'][2])
        self.assertIn('选择训练轮次参与者', rows['LBLS-GN-LBXX-F0064']['cells'][2])
        self.assertIn('不单独追加业务事件', rows['LBLS-GN-LBXX-F0064']['cells'][3])
        self.assertIn('外部能力适配器', rows['LBLS-GN-LBXX-F0028']['cells'][2])
        self.assertTrue(all(row['id'].startswith('FL-SD-OP-F') for row in rows.values()))

    def test_chapter_six_traces_every_current_function_in_both_directions(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        traces = merge.traceability_rows(source)
        self.assertEqual(len(traces), 192)
        self.assertEqual(len({row['id'] for row in traces}), len(traces))
        selected = next(row for row in traces if row['id'] == 'LBLS-GN-LBXX-F0066')
        self.assertIn('FL-SD-OP-F0066', selected['design'])
        self.assertIn('FL-SD-DET-05-R01', selected['design'])
        self.assertIn('FL-SD-OP-F0001', traces[0]['design'])
        self.assertIn('FL-SD-OP-F0192', traces[-1]['design'])
        self.assertNotIn('FL-SU-', traces[0]['design'])

        document = Document(merge.DEFAULT_TARGET)
        merge.update_traceability_chapter(document, traces)
        forward, reverse = merge.traceability_tables(document)
        self.assertEqual(len(forward.rows), len(traces) + 26)
        self.assertEqual(len(reverse.rows), len(merge.reverse_traceability_rows(traces)) + 26)
        for table in (forward, reverse):
            self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
            self.assertTrue(all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows))
        for index, row in enumerate(traces, start=1):
            self.assertEqual(forward.rows[index].cells[1].text, row['requirement'])
            self.assertEqual(forward.rows[index].cells[2].text, row['design'])
            self.assertEqual(forward.rows[index].cells[3].text, '')
        reverse_rows = merge.reverse_traceability_rows(traces)
        reverse_ids = [identifier for row in reverse_rows for identifier in row['requirement'].split('、')]
        self.assertEqual(set(reverse_ids), {row['id'] for row in traces})
        self.assertEqual(len([row for row in traces if 'FL-SD-DET-05' in row['design']]), 19)
        for index, row in enumerate(reverse_rows, start=1):
            self.assertEqual(reverse.rows[index].cells[1].text, row['design'])
            self.assertEqual(reverse.rows[index].cells[2].text, row['requirement'])
            self.assertEqual(reverse.rows[index].cells[3].text, '')
        self.assertIn('LBLS-GN-SZR-HOME', forward.rows[193].cells[1].text)
        self.assertIn('LBLS-GN-SZR-HOME', reverse.rows[len(reverse_rows) + 1].cells[2].text)
        self.assertEqual(forward.rows[193].cells[0].text, '193')
        self.assertEqual(reverse.rows[len(reverse_rows) + 1].cells[0].text, str(len(reverse_rows) + 1))
        self.assertIn('操作处理设计行双向对照', '\n'.join(p.text for p in document.paragraphs))

    def test_traceability_mapping_changes_are_tracked(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        original = merge.traceability_signatures(merge.traceability_rows(source))
        changed = deepcopy(source)
        table = next(block for block in changed['sections']['detailed'] if block['id'] == 'FL-SD-DET-23')
        table['rows'][0]['cells'][2] = 'FL-U-EVAL'
        updated = merge.traceability_signatures(merge.traceability_rows(changed))
        self.assertEqual(merge.changed_entries(original, updated), ['LBLS-GN-LBXX-F0066'])

    def test_context_unit_coverage_rejects_new_unmapped_context(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        context_table = next(block for block in source['sections']['detailed'] if block['id'] == 'FL-SD-DET-26')
        self.assertEqual(len(context_table['rows']), 15)
        changed = deepcopy(source)
        table = next(block for block in changed['sections']['detailed'] if block['id'] == 'FL-SD-DET-26')
        table['rows'] = table['rows'][1:]
        with self.assertRaisesRegex(ValueError, 'Context design units do not cover current requirements'):
            merge.traceability_rows(changed)

    def test_detailed_design_sources_cover_key_units(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        signatures = merge.implementation_source_signatures(source)
        self.assertEqual(len(signatures), 8)
        self.assertIn('FL-SD-DET-05-R02', signatures)
        self.assertIn('FL-SD-DET-05-R07', signatures)

    def test_conceptual_figures_are_embedded_and_tracked(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        figures = merge.figure_assets(source)
        document = Document()
        anchor = document.add_paragraph('数据库概念设计', style='Heading 4')
        merge.insert_section(document, anchor, '{{FIGURE:FL-SD-DBC-04}}', figures)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(document.paragraphs[2].text, '图：联邦学习概念数据关系（治理与数据准备）')
        self.assertNotIn('FIGURE:', '\n'.join(paragraph.text for paragraph in document.paragraphs))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'image.png').write_bytes(b'before')
            (root / 'figure.dot').write_text('source', encoding='utf-8')
            fixture = {'sections': {'db-conceptual': [{
                'id': 'FL-SD-DBC-TEST', 'type': 'figure', 'image': 'image.png',
                'source': 'figure.dot', 'caption': '测试', 'widthCm': 12,
            }]}}
            before = merge.structured_section_items(fixture, root)
            (root / 'image.png').write_bytes(b'after')
            after = merge.structured_section_items(fixture, root)
            self.assertEqual(merge.changed_entries(before, after), ['db-conceptual:FL-SD-DBC-TEST'])

    def test_interface_inventory_covers_actual_resource_routes(self):
        translations = json.loads(merge.MODEL.read_text(encoding='utf-8'))['translations']['zh-CN']
        records = merge.discover_routes(merge.ROOT, translations)
        routes = {(item['provider'], item['method'], item['path']) for item in records}
        self.assertEqual(len(routes), len(records))
        self.assertGreaterEqual(len(records), 320)
        self.assertIn(('共享支撑服务', 'POST', '/uploadedfile/uploadfile/file'), routes)
        self.assertIn(('训练引擎', 'GET', '/jobs/{jobId}/artifacts/{artifactName}'), routes)
        rendered = merge.render_inventory(records)
        self.assertEqual(sum(rendered.count(f'\n| {prefix}-') for prefix in 'PASE'), len(records))
        self.assertIn('上传文件', rendered)
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        detail_rows = next(block['rows'] for block in source['sections']['interfaces'] if block['id'] == 'FL-SD-INT-05')
        inventory_ids = {record['id'] for record in records}
        for row in detail_rows:
            references = re.findall(r'\b[APSE]-[0-9A-F]{8}\b', row['cells'][0])
            self.assertTrue(references, row['id'])
            self.assertTrue(set(references) <= inventory_ids, row['id'])

    def test_table_row_change_reports_stable_id(self):
        source = merge.read_structured_sections(merge.DEFAULT_SECTIONS)
        changed = deepcopy(source)
        table = next(block for block in changed['sections']['architecture'] if block['id'] == 'FL-SD-ARC-02')
        table['rows'][0]['cells'][1] = '调整后的职责'
        original_items = merge.structured_section_items(source)
        changed_items = merge.structured_section_items(changed)
        self.assertEqual(
            merge.changed_entries(original_items, changed_items),
            ['architecture:FL-SD-ARC-02-R01'],
        )

    def test_interface_source_change_reports_affected_row(self):
        source = {'sections': {'interfaces': [{
            'type': 'table', 'rows': [{'id': 'FL-SD-INT-05-R01', 'sourceRefs': ['client.kt']}],
        }]}}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'client.kt'
            path.write_text('before', encoding='utf-8')
            before = merge.interface_source_signatures(source, Path(directory))
            path.write_text('after', encoding='utf-8')
            after = merge.interface_source_signatures(source, Path(directory))
            self.assertEqual(merge.changed_entries(before, after), ['FL-SD-INT-05-R01'])
            messages = merge.register_changes(
                {'interfaceSourceRefs': before},
                {'interfaceSourceRefs': after, 'sourceSections': {}, 'targetSections': {},
                 'targetSha256': None, 'target': None, 'source': None},
            )
            self.assertTrue(any('接口实现来源变更 1 处：FL-SD-INT-05-R01' in message for message in messages))
            path.unlink()
            with self.assertRaisesRegex(ValueError, 'Interface source is missing'):
                merge.interface_source_signatures(source, Path(directory))

    def test_input_output_decisions_render_from_tracked_items(self):
        decisions = merge.read_decisions(merge.DEFAULT_DECISIONS)
        markdown = merge.render_decisions(decisions)
        self.assertEqual(markdown.splitlines().count('| --- | --- | --- | --- |'), 1)
        self.assertEqual(markdown.splitlines().count('| --- | --- | --- | --- | --- |'), 1)
        self.assertEqual(markdown.count('\n| '), len(decisions['items']) + len(decisions['behaviorItems']) + 4)
        self.assertIn('原始样本留在参与方', markdown)
        self.assertNotIn('FL-IO-01', markdown)
        self.assertIn('平台到训练引擎的模型聚合', markdown)
        self.assertIn('联邦学习业务行为与查询展示决策', markdown)

    def test_reviewed_corrections_preserve_original_and_fix_merged_copy(self):
        document = Document(merge.DEFAULT_TARGET)
        source = merge.read_corrections(merge.DEFAULT_CORRECTIONS)
        targets = merge.correction_targets(document, source)
        self.assertEqual(len(targets), len(source['paragraphs']) + len(source['tableCells']) + len(source['referenceRows']))
        self.assertEqual(document.tables[0].rows[4].cells[0].text, '软件需求规格说明')
        merge.apply_corrections(source, targets)
        self.assertEqual(document.tables[0].rows[4].cells[0].text, '软件设计说明')
        self.assertEqual(document.tables[2].rows[-1].cells[2].text, '跨域健康管理联邦孪生平台软件需求规格说明')
        self.assertTrue(any('不以 Nacos、Redis 或共用 PostgreSQL' in p.text for p in document.paragraphs))
        self.assertTrue(any('管理员引导流程' in p.text for p in document.paragraphs))
        self.assertEqual(len(merge.correction_signatures(source)), len(targets))
        changed = deepcopy(source)
        changed['paragraphs'][0]['expected'] = '已变化的总稿表述'
        with self.assertRaisesRegex(ValueError, 'Expected one paragraph for correction'):
            merge.correction_targets(Document(merge.DEFAULT_TARGET), changed)

    def test_behavior_and_context_tables_fit_joint_style(self):
        document = Document(merge.DEFAULT_TARGET)
        cases = (
            (('业务过程', '触发与输入', '状态与校验决策', '结果及异常', '查询展示'),
             [1250, 1600, 2400, 2250, 1500]),
            (('上下文软件单元', '输入与主要判断', '状态输出与异常', '读侧或部署边界'),
             [1700, 2550, 2550, 2200]),
        )
        for headers, expected_widths in cases:
            table = document.add_table(rows=2, cols=len(headers))
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            merge.format_inserted_tables(document, [table])
            self.assertEqual([column.width.twips for column in table.columns], expected_widths)
            self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))

    def test_model_reference_change_identifies_affected_decision(self):
        decisions = {'items': [{'id': 'FL-IO-01', 'commandRefs': ['Example'], 'eventRefs': []}]}
        before = {'slices': [{'commands': [{'name': 'Example', 'fields': [{'name': 'id'}]}]}]}
        after = {'slices': [{'commands': [{'name': 'Example', 'fields': [{'name': 'id'}, {'name': 'value'}]}]}]}
        original = merge.decision_model_signatures(decisions, before)
        changed = merge.decision_model_signatures(decisions, after)
        self.assertEqual(merge.changed_entries(original, changed), ['FL-IO-01'])
        messages = merge.register_changes(
            {'decisionItems': {'FL-IO-01': 'same'}, 'decisionModelRefs': original,
             'sourceSections': {}, 'targetSections': {}, 'targetSha256': 'same',
             'source': 'source.md', 'target': 'target.docx'},
            {'decisionItems': {'FL-IO-01': 'same'}, 'decisionModelRefs': changed,
             'sourceSections': {}, 'targetSections': {}, 'targetSha256': 'same',
             'source': 'source.md', 'target': 'target.docx'},
        )
        self.assertTrue(any('关联建模定义变更 1 处：FL-IO-01' in message for message in messages))
        with self.assertRaisesRegex(ValueError, 'Unknown commands reference in FL-IO-01'):
            merge.decision_model_signatures(decisions, {'slices': []})

    def test_inserted_table_uses_joint_document_style(self):
        document = Document(merge.DEFAULT_TARGET)
        original_count = len(document.tables)
        table = document.add_table(rows=2, cols=3)
        for cell, text in zip(table.rows[0].cells, ('输入或条件', '处理决策', '输出')):
            cell.text = text
        table.rows[1].cells[0].text = '测试输入'
        merge.format_inserted_tables(document, [table])
        self.assertEqual(len(document.tables), original_count + 1)
        self.assertEqual(table.style.name, 'Normal Table')
        self.assertEqual(table.rows[0].cells[0].paragraphs[0].style.name, 'RS-表-表头')
        self.assertEqual(table.rows[1].cells[0].paragraphs[0].style.name, 'RS-表-内容')
        self.assertIsNone(table.rows[0].cells[0]._tc.tcPr.find(qn('w:shd')))
        self.assertAlmostEqual(sum(column.width.cm for column in table.columns), 15.88, delta=0.1)

    def test_interface_inventory_table_prioritizes_paths_and_repeats_header(self):
        document = Document(merge.DEFAULT_TARGET)
        table = document.add_table(rows=2, cols=4)
        for cell, text in zip(table.rows[0].cells, ('接口标识', '调用方', '方法与相对路径', '用途')):
            cell.text = text
        merge.format_inserted_tables(document, [table])
        widths = [column.width.twips for column in table.columns]
        self.assertEqual(widths, [1500, 1500, 4500, 1500])
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
        self.assertTrue(all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows))

    def test_detailed_design_table_gives_explanation_room(self):
        document = Document(merge.DEFAULT_TARGET)
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = '设计要素'
        table.rows[0].cells[1].text = '设计说明'
        merge.format_inserted_tables(document, [table])
        self.assertEqual([column.width.twips for column in table.columns], [2000, 7000])
        self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
        self.assertTrue(all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows))

    def test_performance_sections_do_not_mingle(self):
        document = Document()
        performance = document.add_paragraph('性能设计', style='Heading 2')
        document.add_paragraph('其他模块原有性能内容')
        interfaces = document.add_paragraph('接口设计', style='Heading 2')
        merge.insert_section(document, performance, merge.OTHER_MODULES_PERFORMANCE)
        merge.insert_before_heading(document, interfaces, '### 联邦学习模块性能设计\n\n联邦学习性能内容')
        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ['性能设计', '医疗数字人及风险动态评估模块性能设计',
             '其他模块原有性能内容', '联邦学习模块性能设计',
             '联邦学习性能内容', '接口设计'],
        )

    def test_target_sections_locate_changed_heading(self):
        original = Document()
        original.add_paragraph('模块 A', style='Heading 2')
        original.add_paragraph('原内容')
        original.add_paragraph('模块 B', style='Heading 2')
        original.add_paragraph('保持不变')
        changed = Document()
        changed.add_paragraph('模块 A', style='Heading 2')
        changed.add_paragraph('调整后的内容')
        changed.add_paragraph('模块 B', style='Heading 2')
        changed.add_paragraph('保持不变')
        self.assertEqual(
            merge.changed_entries(merge.document_sections(original), merge.document_sections(changed)),
            ['模块 A [1]'],
        )

    def test_source_change_and_format_only_change_are_reported(self):
        baseline = {
            'source': 'source.md', 'target': 'target.docx', 'targetSha256': 'before',
            'sourceSections': {'decisions': 'before'}, 'targetSections': {'section': 'same'},
        }
        current = {
            **baseline, 'targetSha256': 'after',
            'sourceSections': {'decisions': 'after'},
        }
        messages = merge.register_changes(baseline, current)
        self.assertTrue(any('联邦学习正文变更 1 处：decisions' in message for message in messages))
        self.assertTrue(any('章节文字未变' in message for message in messages))


if __name__ == '__main__':
    unittest.main()
