"""Checks for the reviewed 3.3.1 delivery layout."""

import importlib.util
import re
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('format_438c_331', ROOT / 'scripts/format-438c-331.py')
formatter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(formatter)
sys.path.insert(0, str(ROOT / 'docs/document-pack/federation-learning'))
from markdown_to_docx import format_requirements_tables, write_markdown


class RequirementLayoutTest(unittest.TestCase):
    def test_word_section_headings_use_final_document_levels(self):
        document = Document()
        write_markdown(document, '# 3.3.1 功能需求\n\n## 3.3.1.2 组织管理', '',
                       heading_style_offset=2, strip_heading_numbers=True)
        self.assertEqual(document.paragraphs[0].style.name, 'Heading 3')
        self.assertEqual(document.paragraphs[1].style.name, 'Heading 4')
        self.assertEqual(document.paragraphs[0].text, '功能需求')
        self.assertEqual(document.paragraphs[1].text, '组织管理')

    def test_word_tables_match_joint_requirements_format(self):
        document = Document()
        headers = [
            ['序号', '一级功能名称/标识', '二级功能名称/标识', '三级功能名称/标识', '备注'],
            ['字段名称', '输入形式', '可修改', '默认值', '内容及规则', '必填'],
            ['输入项', '内容及规则'],
            ['输出项', '内容及规则'],
        ]
        tables = []
        for names in headers:
            table = document.add_table(rows=2, cols=len(names))
            for cell, name in zip(table.rows[0].cells, names):
                cell.text = name
            table.rows[1].cells[0].text = '测试内容'
            tables.append(table)
        format_requirements_tables(document)
        for table in tables:
            self.assertEqual(table.rows[0].cells[0].paragraphs[0].style.name, 'XG-表-头')
            self.assertEqual(table.rows[1].cells[0].paragraphs[0].style.name, 'XG-表')
            self.assertIsNone(table.rows[0].cells[0]._tc.tcPr.find(qn('w:shd')))
            border = table.rows[0].cells[0]._tc.tcPr.find(qn('w:tcBorders'))
            self.assertEqual(border.find(qn('w:top')).get(qn('w:color')), '000000')
            self.assertTrue(table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None)
        self.assertAlmostEqual(sum(column.width.cm for column in tables[0].columns), 16, delta=0.1)

    def test_manual_operation_keeps_inputs_constraints_and_results(self):
        result = formatter.format_function('注册组织', 'LBLS-GN-LBXX-F0001', [
            '系统应支持注册组织，由组织管理员执行。',
            '输入要求：', '- 组织名称（文本，必填）。',
            '业务约束：', '- 组织名称必须唯一。',
            '成功结果：', '- 组织已注册。',
            '不满足约束时，应拒绝执行并说明原因，包括：组织已存在。',
        ])
        self.assertIn('**用户故事：**作为组织管理员', result)
        self.assertIn('**需求说明**\n\n系统应支持注册组织，由组织管理员执行。组织名称必须唯一。完成后应形成“组织已注册”的业务结果。', result)
        self.assertIn('**1）用户输入**\n\n**表 4：用户输入字段**', result)
        self.assertIn('| 字段名称 | 输入形式 | 可修改 | 默认值 | 内容及规则 | 必填 |', result)
        self.assertIn('| 组织名称 | 文本输入 | 是 | — | 类型：文本。 | 是 |', result)
        self.assertIn('**2）处理逻辑**', result)
        self.assertIn('业务约束：组织名称必须唯一', result)
        self.assertIn('**3）异常处理**', result)
        self.assertIn('组织已存在', result)
        self.assertIn('**（3）输出**', result)
        self.assertIn('**表 5：输出信息**', result)
        self.assertIn('| 成功结果 | 组织已注册。 |', result)

    def test_automated_operation_does_not_call_system_values_user_input(self):
        result = formatter.format_function('使文件过期', 'LBLS-GN-LBXX-F0159', [
            '系统应支持使文件过期，由业务过程触发。',
            '自动触发条件：保留期过后使文件过期。',
            '输入要求：', '- 过期时间（日期时间，必填）。',
            '成功结果：', '- 文件已过期。',
        ])
        self.assertIn('不要求用户填写输入', result)
        self.assertIn('**需求说明**\n\n系统应在保留期过后自动使文件过期。完成后应形成“文件已过期”的业务结果。', result)
        self.assertIn('**2）系统输入与触发**\n\n**表 5：系统输入与触发**', result)
        self.assertIn('| 触发条件 | 保留期过后使文件过期。 |', result)
        self.assertIn('| 系统输入 | 过期时间（日期时间，必填）。 |', result)
        self.assertNotIn('自动使文件过期时自动使文件过期', result)

    def test_query_retains_filter_and_display_requirements(self):
        result = formatter.format_function('组织目录', 'LBLS-GN-LBXX-F0005', [
            '系统应提供组织目录，供相关用户查询。',
            '查询与展示：', '- 应提供组织目录，包含组织 ID、组织名称、状态。',
            '- 应支持按组织 ID 精确定位，并按状态筛选。',
        ])
        self.assertIn('**1）用户输入**', result)
        self.assertIn('**需求说明**\n\n系统应提供组织目录，展示授权范围内的业务信息。应支持按组织 ID 精确定位，并按状态筛选。', result)
        self.assertIn('| 查询条件 | 条件选择 | 是 | — | 用户在授权范围内选择页面提供的查询条件', result)
        self.assertIn('稳定排序分页', result)
        self.assertIn('无匹配记录时提示无结果', result)
        self.assertIn('按组织 ID 精确定位，并按状态筛选', result)

    def test_explicit_field_source_is_rendered_without_guessing_other_sources(self):
        fields = [
            {'name': 'organizationName', 'source': {'kind': 'derived', 'from': ['OrganizationDirectory.organizationName']}},
            {'name': 'activationNote'},
        ]
        translations = {
            'Organization Name': '组织名称', 'Organization Directory': '组织目录',
            'Activation Note': '激活备注',
        }
        result = formatter.format_function('激活组织', 'LBLS-GN-LBXX-F0002', [
            '系统应支持激活组织，由平台管理员执行。',
            '输入要求：', '- 组织名称（文本，必填）。', '- 激活备注（文本，可选）。',
        ], source_fields=fields, translations=translations)
        self.assertIn('| 组织名称 | 文本输入 | 是 | — | 类型：文本；来源：组织目录中的组织名称。 | 是 |', result)
        self.assertIn('| 激活备注 | 文本输入 | 是 | — | 类型：文本。 | 否 |', result)

    def test_event_source_is_named_as_an_event(self):
        field = {'name': 'organizationId', 'source': {'kind': 'direct', 'from': ['PlanCreated.organizationId']}}
        description = formatter.source_description(field, {
            'Plan Created': '计划已创建', 'Organization Id': '组织 ID',
        }, {'PlanCreated'})
        self.assertEqual(description, '计划已创建事件中的组织 ID')

    def test_all_reviewed_functions_are_preserved(self):
        source = (ROOT / 'docs/438c/merge/3.3.1-federation-learning.md').read_text(encoding='utf-8')
        result = formatter.format_document(source)
        source_ids = set(re.findall(r'^\*\*.+? / (LBLS-GN-LBXX-F\d{4})\*\*$', source, re.MULTILINE))
        function_count = len(source_ids)
        self.assertEqual(function_count, 192)
        self.assertEqual(result.count('**用户故事：**'), function_count)
        self.assertEqual(result.count('**需求说明**'), function_count)
        self.assertNotIn('**需求说明**\n\n**（1）输入**', result)
        self.assertEqual(result.count('**（1）输入**'), function_count)
        self.assertEqual(result.count('**（2）处理**'), function_count)
        self.assertEqual(result.count('**（3）输出**'), function_count)
        self.assertEqual(result.count('| 输出项 | 内容及规则 |'), function_count)
        self.assertEqual(result.count('| 字段名称 | 输入形式 | 可修改 | 默认值 | 内容及规则 | 必填 |'), function_count)
        table_numbers = [int(value) for value in re.findall(r'^\*\*表 (\d+)：', result, re.MULTILINE)]
        self.assertEqual(table_numbers, list(range(4, 4 + len(table_numbers))))
        result_ids = set(re.findall(r'^\*\*.+? / (LBLS-GN-LBXX-F\d{4})\*\*$', result, re.MULTILINE))
        self.assertEqual(source_ids, result_ids)
        self.assertIn('列表总条数应随筛选条件变化', result)
        self.assertIn('来源：组织目录中的组织名称', result)


if __name__ == '__main__':
    unittest.main()
