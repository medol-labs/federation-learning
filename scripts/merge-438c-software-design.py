#!/usr/bin/env python3
"""Insert reviewed Federation Learning design text into a copy of the joint SDD."""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TARGET = (
    ROOT.parent
    / 'tmp/merge/02“感一通一算”一体化的跨域健康管理联邦孪生平台-软件设计说明v1.0.docx'
)
DEFAULT_SOURCE = ROOT / 'docs/438c/merge/software-design-federation-learning.md'
DEFAULT_DECISIONS = ROOT / 'docs/438c/merge/software-design-input-output.json'
DEFAULT_SECTIONS = ROOT / 'docs/438c/merge/software-design-sections.json'
DEFAULT_CORRECTIONS = ROOT / 'docs/438c/merge/software-design-corrections.json'
DEFAULT_OUTPUT = ROOT / 'docs/generated/438c/merge/software-design-federation-learning-merged.docx'
REGISTER = ROOT / 'docs/438c/merge/software-design-register.json'
MODEL = ROOT / '.medol/codegen-model.json'
REQUIREMENT_REGISTER = ROOT / 'docs/438c/merge/3.3.1-function-register.json'
REQUIREMENT_TEXT = ROOT / 'docs/438c/merge/3.3.1-federation-learning.md'
sys.path.insert(0, str(ROOT / 'docs/document-pack/federation-learning'))
sys.path.insert(0, str(ROOT / 'scripts'))
from markdown_to_docx import write_markdown  # noqa: E402
from software_design_interface_inventory import discover_routes, render_inventory  # noqa: E402
from software_design_operation_inventory import operation_rows, render_operation_design  # noqa: E402


ANCHORS = {
    'decisions': ('联邦学习模块输入/输出设计决策/LBLS-GN-LBXX', 'Heading 3'),
    'db-conceptual': ('联邦学习', 'Heading 5'),
    'db-logical': ('联邦学习', 'Heading 5'),
    'db-physical': ('联邦学习', 'Heading 5'),
    'architecture': ('联邦学习模块CSCI部件', 'Heading 3'),
    'execution': ('联邦学习模块执行方案', 'Heading 3'),
    'performance': ('性能设计', 'Heading 2'),
    'interfaces': ('接口设计', 'Heading 2'),
    'detailed': ('联邦学习模块CSCI详细设计', 'Heading 2'),
}
DATABASE_PARENTS = {
    'db-conceptual': '数据库概念设计',
    'db-logical': '数据库逻辑设计',
    'db-physical': '数据库物理设计',
}
MARKER = re.compile(r'^<!-- INSERT ([a-z-]+) -->$', re.MULTILINE)
DECISIONS_PLACEHOLDER = '{{INPUT_OUTPUT_DECISIONS}}'
SECTION_PLACEHOLDER = '{{SOFTWARE_DESIGN_SECTION}}'
FIGURE_MARKER = re.compile(r'^\{\{FIGURE:([A-Za-z0-9-]+)\}\}$')
OTHER_MODULES_PERFORMANCE = '### 医疗数字人及风险动态评估模块性能设计'


def parse_sections(text):
    markers = list(MARKER.finditer(text))
    if set(match.group(1) for match in markers) != set(ANCHORS):
        raise ValueError('Design merge source must contain each expected INSERT marker exactly once.')
    if len(markers) != len(ANCHORS):
        raise ValueError('Duplicate design merge marker.')
    return {
        match.group(1): text[match.end():markers[index + 1].start() if index + 1 < len(markers) else len(text)].strip()
        for index, match in enumerate(markers)
    }


def read_decisions(path):
    decisions = json.loads(path.read_text(encoding='utf-8'))
    if decisions.get('schemaVersion') != 1:
        raise ValueError('Unsupported input/output design source version.')
    if not decisions.get('introduction') or not decisions.get('conclusion'):
        raise ValueError('Input/output design source requires introduction and conclusion.')
    items = decisions.get('items', [])
    behavior_items = decisions.get('behaviorItems', [])
    ids = [item.get('id') for item in (*items, *behavior_items)]
    if not ids or any(not id_ for id_ in ids) or len(set(ids)) != len(ids):
        raise ValueError('Input/output design items require unique stable IDs.')
    for item in items:
        if any(not item.get(key) for key in ('direction', 'input', 'output', 'decision')):
            raise ValueError(f'Input/output design item is incomplete: {item["id"]}')
        if not item.get('commandRefs') and not item.get('eventRefs'):
            raise ValueError(f'Input/output design item has no model references: {item["id"]}')
    for item in behavior_items:
        if any(not item.get(key) for key in ('scenario', 'trigger', 'decision', 'outcome', 'display')):
            raise ValueError(f'Behavior design item is incomplete: {item["id"]}')
        if not item.get('commandRefs') and not item.get('eventRefs'):
            raise ValueError(f'Behavior design item has no model references: {item["id"]}')
    return decisions


def read_corrections(path):
    source = json.loads(path.read_text(encoding='utf-8'))
    if source.get('schemaVersion') != 1:
        raise ValueError('Unsupported software design corrections source version.')
    ids = set()
    for kind in ('paragraphs', 'tableCells'):
        for item in source.get(kind, []):
            if (not item.get('id') or item['id'] in ids or not item.get('expected')
                    or not item.get('replacement') or item['expected'] == item['replacement']):
                raise ValueError(f'Invalid software design correction: {item.get("id")}')
            ids.add(item['id'])
    for item in source.get('referenceRows', []):
        if (not item.get('id') or item['id'] in ids or not item.get('headers')
                or len(item.get('cells', [])) != len(item['headers'])):
            raise ValueError(f'Invalid software design reference row: {item.get("id")}')
        ids.add(item['id'])
    return source


def correction_signatures(source):
    return {
        item['id']: digest(json.dumps(item, ensure_ascii=False, sort_keys=True))
        for kind in ('paragraphs', 'tableCells', 'referenceRows') for item in source.get(kind, [])
    }


def correction_targets(document, source):
    targets = {}
    for item in source.get('paragraphs', []):
        matches = [p for p in document.paragraphs if p.text.strip() == item['expected']]
        if len(matches) != 1:
            raise ValueError(f'Expected one paragraph for correction {item["id"]}, found {len(matches)}.')
        targets[item['id']] = matches[0]
    for item in source.get('tableCells', []):
        matches = {
            id(cell._tc): cell
            for table in document.tables for row in table.rows for cell in row.cells
            if cell.text.strip() == item['expected']
        }
        if len(matches) != 1:
            raise ValueError(f'Expected one table cell for correction {item["id"]}, found {len(matches)}.')
        targets[item['id']] = next(iter(matches.values()))
    for item in source.get('referenceRows', []):
        matches = [
            table for table in document.tables
            if [cell.text.strip() for cell in table.rows[0].cells] == item['headers']
        ]
        if len(matches) != 1 or any(
            [cell.text.strip() for cell in row.cells] == item['cells']
            for row in matches[0].rows
        ):
            raise ValueError(f'Expected one unchanged reference table for {item["id"]}.')
        targets[item['id']] = matches[0]
    return targets


def apply_corrections(source, targets):
    for item in source.get('paragraphs', []):
        paragraph = targets[item['id']]
        paragraph.text = item['replacement']
    for item in source.get('tableCells', []):
        cell = targets[item['id']]
        runs = [run for paragraph in cell.paragraphs for run in paragraph.runs if item['expected'] in run.text]
        if len(runs) != 1:
            raise ValueError(f'Cannot preserve formatting for correction {item["id"]}.')
        runs[0].text = runs[0].text.replace(item['expected'], item['replacement'])
    for item in source.get('referenceRows', []):
        table = targets[item['id']]
        new_row = deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_row)
        for cell, value in zip(_Row(new_row, table).cells, item['cells']):
            cell.text = value


def render_decisions(decisions):
    def cell(value):
        return value.replace('|', '\\|').replace('\n', ' ')

    rows = [
        decisions['introduction'],
        '',
        '| 交换方向 | 主要输入 | 主要输出 | 输入/输出设计决策 |',
        '| --- | --- | --- | --- |',
    ]
    for item in decisions['items']:
        rows.append('| ' + ' | '.join(cell(item[key]) for key in ('direction', 'input', 'output', 'decision')) + ' |')
    rows.extend(('', decisions['conclusion']))
    if decisions.get('behaviorItems'):
        rows.extend(('', '#### 联邦学习业务行为与查询展示决策', '',
                     '| 业务过程 | 触发与输入 | 状态与校验决策 | 结果及异常 | 查询展示 |',
                     '| --- | --- | --- | --- | --- |'))
        for item in decisions['behaviorItems']:
            rows.append('| ' + ' | '.join(cell(item[key]) for key in (
                'scenario', 'trigger', 'decision', 'outcome', 'display'
            )) + ' |')
    return '\n'.join(rows)


def read_structured_sections(path):
    source = json.loads(path.read_text(encoding='utf-8'))
    if source.get('schemaVersion') != 1 or set(source.get('sections', {})) != set(ANCHORS) - {'decisions'}:
        raise ValueError('Software design source must define the other eight insertion sections.')
    seen = set()
    for key, blocks in source['sections'].items():
        if not blocks:
            raise ValueError(f'Software design section is empty: {key}')
        for block in blocks:
            id_ = block.get('id')
            if not id_ or id_ in seen:
                raise ValueError(f'Duplicate or missing software design item ID: {id_}')
            seen.add(id_)
            kind = block.get('type')
            if kind == 'table':
                width = len(block.get('headers', []))
                if width < 2 or not block.get('rows'):
                    raise ValueError(f'Software design table is incomplete: {id_}')
                for row in block['rows']:
                    row_id = row.get('id')
                    if not row_id or row_id in seen or len(row.get('cells', [])) != width:
                        raise ValueError(f'Software design table row is invalid: {row_id}')
                    seen.add(row_id)
            elif kind == 'paragraph' and not block.get('text'):
                raise ValueError(f'Software design paragraph is empty: {id_}')
            elif kind == 'heading' and (block.get('level') not in (3, 4, 5) or not block.get('text')):
                raise ValueError(f'Software design heading is invalid: {id_}')
            elif kind == 'readModelInventory' and key != 'db-logical':
                raise ValueError(f'Read model inventory belongs in db-logical: {id_}')
            elif kind == 'interfaceInventory' and key != 'interfaces':
                raise ValueError(f'Interface inventory belongs in interfaces: {id_}')
            elif kind == 'operationDesign' and key != 'detailed':
                raise ValueError(f'Operation design belongs in detailed design: {id_}')
            elif kind == 'figure' and (
                not block.get('image') or not block.get('source') or not block.get('caption')
                or not isinstance(block.get('widthCm'), (int, float))
                or not 0 < block['widthCm'] <= 16.5
            ):
                raise ValueError(f'Software design figure is incomplete: {id_}')
            elif kind not in ('table', 'paragraph', 'heading', 'readModelInventory',
                              'interfaceInventory', 'operationDesign', 'figure'):
                raise ValueError(f'Unsupported software design item type: {kind}')
    return source


def render_structured_sections(source, operation_groups=None):
    def row(cells):
        return '| ' + ' | '.join(value.replace('|', '\\|').replace('\n', ' ') for value in cells) + ' |'

    rendered = {}
    for key, blocks in source['sections'].items():
        fragments = []
        for block in blocks:
            kind = block['type']
            if kind == 'paragraph':
                fragments.append(block['text'])
            elif kind == 'heading':
                fragments.append('#' * block['level'] + ' ' + block['text'])
            elif kind == 'readModelInventory':
                fragments.append('{{READ_MODEL_INVENTORY}}')
            elif kind == 'interfaceInventory':
                fragments.append('{{INTERFACE_INVENTORY}}')
            elif kind == 'operationDesign':
                if operation_groups is None:
                    operation_groups = load_operation_design()
                fragments.append(render_operation_design(operation_groups))
            elif kind == 'figure':
                fragments.append(f'{{{{FIGURE:{block["id"]}}}}}')
            else:
                headers = block['headers']
                lines = [row(headers), row(['---'] * len(headers))]
                for item in block['rows']:
                    cells = list(item['cells'])
                    if block['id'] in ('FL-SD-DET-05', 'FL-SD-DET-26'):
                        cells[0] = f'{cells[0]} / {item["id"]}'
                    lines.append(row(cells))
                fragments.append('\n'.join(lines))
        rendered[key] = '\n\n'.join(fragments)
    return rendered


def figure_asset_path(root, reference):
    root = root.resolve()
    path = (root / reference).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError(f'Software design figure asset is missing or outside the project: {reference}')
    return path


def figure_assets(source, root=ROOT):
    return {
        block['id']: (block, figure_asset_path(root, block['image']))
        for blocks in source['sections'].values()
        for block in blocks if block['type'] == 'figure'
    }


def structured_section_items(source, root=ROOT):
    items = {}
    for key, blocks in source['sections'].items():
        for block in blocks:
            # Table rows get independent identities so one edited row is reported precisely.
            summary = {field: value for field, value in block.items() if field != 'rows'}
            if block['type'] == 'figure':
                for field in ('image', 'source'):
                    summary[f'{field}Sha256'] = hashlib.sha256(
                        figure_asset_path(root, block[field]).read_bytes()
                    ).hexdigest()
            items[f'{key}:{block["id"]}'] = digest(json.dumps(summary, ensure_ascii=False, sort_keys=True))
            for row in block.get('rows', []):
                items[f'{key}:{row["id"]}'] = digest(json.dumps(row, ensure_ascii=False, sort_keys=True))
    return items


def requirement_signatures(source, register_path=REQUIREMENT_REGISTER):
    register = json.loads(register_path.read_text(encoding='utf-8'))
    functions = {item['id']: item for item in register['functions']}
    trace_tables = [
        block for block in source['sections']['detailed']
        if block['id'] == 'FL-SD-DET-23' and block['type'] == 'table'
    ]
    if len(trace_tables) != 1:
        raise ValueError('Detailed design must contain one functional requirement trace table.')
    unit_tables = [
        block for block in source['sections']['detailed']
        if block['id'] == 'FL-SD-DET-05' and block['type'] == 'table'
    ]
    if len(unit_tables) != 1:
        raise ValueError('Detailed design must contain one software unit inventory.')
    known_units = {row['cells'][0].split(' ', 1)[0] for row in unit_tables[0]['rows']}
    mapped_units = set()
    signatures = {}
    for row in trace_tables[0]['rows']:
        requirement_id, title, units = row['cells']
        requirement = functions.get(requirement_id)
        if requirement is None or requirement['title'] != title or not units:
            raise ValueError(f'Detailed design requirement trace is invalid: {row["id"]}')
        row_units = set(units.split('、'))
        if not row_units <= known_units:
            raise ValueError(f'Detailed design requirement trace has unknown unit: {row["id"]}')
        mapped_units.update(row_units)
        if requirement_id in signatures:
            raise ValueError(f'Duplicate detailed design requirement trace: {requirement_id}')
        signatures[requirement_id] = digest(json.dumps(requirement, ensure_ascii=False, sort_keys=True))
    if mapped_units != known_units:
        raise ValueError(f'Detailed design units without requirement traces: {known_units - mapped_units}')
    return signatures


def load_operation_design(model_path=MODEL, register_path=REQUIREMENT_REGISTER,
                          requirement_text_path=REQUIREMENT_TEXT):
    return operation_rows(
        json.loads(model_path.read_text(encoding='utf-8')),
        json.loads(register_path.read_text(encoding='utf-8')),
        requirement_text_path.read_text(encoding='utf-8'),
    )


def traceability_rows(source, register_path=REQUIREMENT_REGISTER, operation_groups=None):
    """Trace requirement IDs to actual design items in this document."""
    register = json.loads(register_path.read_text(encoding='utf-8'))
    if len([block for block in source['sections']['detailed'] if block['type'] == 'operationDesign']) != 1:
        raise ValueError('Detailed design must contain one operation design inventory.')
    context_table = next(block for block in source['sections']['detailed'] if block['id'] == 'FL-SD-DET-26')
    context_design_units = {}
    for item in context_table['rows']:
        name = item.get('sourceContext')
        if not name or name in context_design_units:
            raise ValueError(f'Duplicate or missing context design unit: {item["id"]}')
        context_design_units[name] = item
    required_contexts = {item['sourceContext'] for item in register['functions']}
    if set(context_design_units) != required_contexts:
        raise ValueError(f'Context design units do not cover current requirements: '
                         f'{required_contexts - set(context_design_units)}')

    if operation_groups is None:
        operation_groups = load_operation_design(register_path=register_path)
    operation_design = {row['requirementId']: row for group in operation_groups for row in group['rows']}
    if len(operation_design) != sum(len(group['rows']) for group in operation_groups):
        raise ValueError('Operation design contains duplicate requirement IDs.')
    if set(operation_design) != {item['id'] for item in register['functions']}:
        raise ValueError('Operation design does not cover the current SRS function register.')

    mapping_table = next(block for block in source['sections']['detailed'] if block['id'] == 'FL-SD-DET-23')
    key_table = next(block for block in source['sections']['detailed'] if block['id'] == 'FL-SD-DET-05')
    key_units = {row['cells'][0].split(' ', 1)[0]: row for row in key_table['rows']}
    detailed_by_requirement = {}
    for mapping in mapping_table['rows']:
        requirement_id, title, unit_codes = mapping['cells']
        if requirement_id in detailed_by_requirement:
            raise ValueError(f'Duplicate detailed design mapping: {requirement_id}')
        codes = unit_codes.split('、')
        if any(code not in key_units for code in codes):
            raise ValueError(f'Unknown key design item: {mapping["id"]}')
        detailed_by_requirement[requirement_id] = (title, [key_units[code] for code in codes])
    rows = []
    seen = set()
    for requirement in register['functions']:
        identifier = requirement['id']
        if identifier in seen:
            raise ValueError(f'Duplicate function requirement: {identifier}')
        seen.add(identifier)
        operation = operation_design[identifier]
        if operation['title'] != requirement['title']:
            raise ValueError(f'Operation design title differs: {identifier}')
        design_refs = [f'第5章 操作处理设计/{operation["id"]}（{operation["title"]}）']
        mapping = detailed_by_requirement.get(identifier)
        if mapping:
            if mapping[0] != requirement['title']:
                raise ValueError(f'Detailed design requirement title differs: {identifier}')
            design_refs.extend(
                f'第5章 关键软件单元及处理决策/{item["id"]}（{item["cells"][0]}）'
                for item in mapping[1]
            )
        rows.append({
            'id': identifier,
            'requirement': f'{requirement["title"]}/{identifier}',
            'design': '；'.join(design_refs),
            'designRefs': design_refs,
            'sourceHash': requirement['sourceHash'],
        })
    if set(detailed_by_requirement) - seen:
        raise ValueError('Detailed design traces contain requirements absent from the current register.')
    return rows


def reverse_traceability_rows(rows):
    """Collect each design item once and list the requirements it addresses."""
    by_design = {}
    for row in rows:
        for design in row['designRefs']:
            entry = by_design.setdefault(design, {'design': design, 'requirements': []})
            entry['requirements'].append(row['id'])
    return [
        {'design': entry['design'], 'requirement': '、'.join(entry['requirements'][index:index + 12])}
        for entry in by_design.values()
        for index in range(0, len(entry['requirements']), 12)
    ]


def traceability_signatures(rows):
    return {
        row['id']: digest(json.dumps(row, ensure_ascii=False, sort_keys=True)) for row in rows
    }


def section_source_signatures(source, section, root=ROOT):
    signatures = {}
    root = root.resolve()
    for block in source['sections'][section]:
        if block['type'] != 'table':
            continue
        for row in block['rows']:
            refs = row.get('sourceRefs', [])
            if not refs:
                continue
            if not isinstance(refs, list) or any(not isinstance(ref, str) or not ref for ref in refs):
                raise ValueError(f'Invalid interface source references: {row["id"]}')
            entries = []
            for ref in refs:
                path = (root / ref).resolve()
                if not path.is_relative_to(root) or not path.is_file():
                    raise ValueError(f'Interface source is missing or outside the project: {row["id"]}: {ref}')
                entries.append((ref, hashlib.sha256(path.read_bytes()).hexdigest()))
            signatures[row['id']] = digest(json.dumps(entries, ensure_ascii=False))
    return signatures


def interface_source_signatures(source, root=ROOT):
    return section_source_signatures(source, 'interfaces', root)


def implementation_source_signatures(source, root=ROOT):
    return section_source_signatures(source, 'detailed', root)


def decision_model_signatures(decisions, model):
    indexed = {
        kind: {entry['name']: entry for slice_ in model['slices'] for entry in slice_.get(kind, [])}
        for kind in ('commands', 'events')
    }
    signatures = {}
    for item in (*decisions['items'], *decisions.get('behaviorItems', [])):
        references = []
        for kind, key in (('commands', 'commandRefs'), ('events', 'eventRefs')):
            for name in item.get(key, []):
                entry = indexed[kind].get(name)
                if entry is None:
                    raise ValueError(f'Unknown {kind} reference in {item["id"]}: {name}')
                references.append((kind, name, entry))
        signatures[item['id']] = digest(json.dumps(references, ensure_ascii=False, sort_keys=True))
    return signatures


def read_model_inventory():
    model = json.loads(MODEL.read_text(encoding='utf-8'))
    translations = model.get('translations', {}).get('zh-CN', {})
    deployment_labels = {
        'FederationLearningPlatform': '平台',
        'FederationLearningRuntimeAgent': '参与方代理',
        'FederationLearningSupport': '共享支撑服务',
    }
    contexts = {
        context['name']: deployment_labels[deployment['name']]
        for deployment in model['deployments']
        if deployment['name'] in deployment_labels
        for context in deployment['contexts']
    }
    rows = ['| 部署单元 | 数据主题与查询视图 | 标识字段 |', '| --- | --- | --- |']
    seen = set()
    signatures = {}
    for slice_ in model['slices']:
        context = slice_['context']
        if context not in contexts:
            continue
        context_title = next(item['title'] for item in model['contexts'] if item['name'] == context)
        for read_model in slice_.get('readmodels', []):
            if read_model['name'] in seen:
                raise ValueError(f"Duplicate read model: {read_model['name']}")
            seen.add(read_model['name'])
            signatures[read_model['name']] = digest(json.dumps(read_model, ensure_ascii=False, sort_keys=True))
            ids = [field['name'] for field in read_model.get('fields', []) if field.get('idAttribute')]
            values = [
                contexts[context],
                f"{translations.get(context_title, context_title)}：{translations.get(read_model['title'], read_model['title'])}",
                '、'.join(ids) if ids else '按持久化实体定义',
            ]
            rows.append('| ' + ' | '.join(value.replace('|', '\\|') for value in values) + ' |')
    if len(seen) < 20:
        raise ValueError('Read model inventory unexpectedly small; check current model export.')
    return '\n'.join(rows), signatures


def find_anchors(document):
    matches = {}
    previous_heading4 = None
    database_index = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name == 'Heading 4':
            previous_heading4 = paragraph.text.strip()
        text = paragraph.text.strip()
        if text == '联邦学习' and paragraph.style.name == 'Heading 5':
            database_index += 1
            key = ('db-conceptual', 'db-logical', 'db-physical')[database_index - 1]
            if previous_heading4 != DATABASE_PARENTS[key]:
                raise ValueError(f'Unexpected database parent for {key}: {previous_heading4}')
            matches[key] = paragraph
            continue
        for key, (expected_text, expected_style) in ANCHORS.items():
            if key.startswith('db-'):
                continue
            if text == expected_text and paragraph.style.name == expected_style:
                if key in matches:
                    raise ValueError(f'Duplicate target anchor: {key}')
                matches[key] = paragraph
    if set(matches) != set(ANCHORS):
        raise ValueError(f'Target outline differs; missing anchors: {set(ANCHORS) - set(matches)}')
    return matches


def traceability_tables(document):
    chapter = None
    tables = []
    for element in document.element.body:
        if element.tag == qn('w:p'):
            paragraph = Paragraph(element, document._body)
            if paragraph.style.name == 'Heading 1':
                if paragraph.text.strip() == '需求可追踪性':
                    if chapter is not None:
                        raise ValueError('Duplicate requirement traceability chapter.')
                    chapter = paragraph
                elif chapter is not None:
                    break
        elif chapter is not None and element.tag == qn('w:tbl'):
            tables.append(Table(element, document._body))
    if chapter is None or len(tables) != 2:
        raise ValueError('Joint design must contain the forward and reverse traceability tables.')
    return tables


def update_traceability_table(document, table, rows, reverse=False):
    expected = (
        ('本文档中被追踪内容的名称/标识', '跨域健康管理联邦孪生平台软件需求规格说明文档中被追踪内容的名称/标识')
        if reverse else
        ('跨域健康管理联邦孪生平台软件需求规格说明文档中被追踪内容的名称/标识', '本文档中被追踪内容的名称/标识')
    )
    if len(table.columns) != 4 or tuple(cell.text.strip() for cell in table.rows[0].cells[1:3]) != expected:
        raise ValueError('Joint design traceability table headings have changed.')
    old_rows = [row for row in table.rows[1:] if 'LBLS-GN-LBXX-' in row.cells[2 if reverse else 1].text]
    if len(old_rows) != 10 or table.rows[1]._tr is not old_rows[0]._tr:
        raise ValueError('Expected ten original Federation Learning traceability placeholders.')
    if any(row.cells[1 if reverse else 2].text.strip() for row in old_rows):
        raise ValueError('Federation Learning traceability placeholders are no longer empty.')
    template = deepcopy(old_rows[0]._tr)
    first_other = table.rows[11]._tr
    for row in old_rows:
        table._tbl.remove(row._tr)
    traces = reverse_traceability_rows(rows) if reverse else rows
    for trace in traces:
        new_row = deepcopy(template)
        first_other.addprevious(new_row)
        cells = _Row(new_row, table).cells
        values = (
            ('', trace['design'], trace['requirement'], '') if reverse else
            ('', trace['requirement'], trace['design'], '')
        )
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.style = document.styles['RS-表-内容']
    for index, row in enumerate(table.rows[1:], start=1):
        row.cells[0].text = str(index)
        for paragraph in row.cells[0].paragraphs:
            paragraph.style = document.styles['RS-表-内容']
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in table.rows:
        if row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is None:
            row._tr.trPr.append(OxmlElement('w:cantSplit'))


def update_traceability_chapter(document, rows):
    forward, reverse = traceability_tables(document)
    update_traceability_table(document, forward, rows)
    update_traceability_table(document, reverse, rows, reverse=True)
    chapter_seen = False
    for paragraph in document.paragraphs:
        if paragraph.style.name == 'Heading 1' and paragraph.text.strip() == '需求可追踪性':
            chapter_seen = True
        elif chapter_seen and paragraph.style.name == 'Heading 2':
            insert_before_heading(
                document, paragraph,
                '联邦学习功能需求按本版软件需求规格说明中的功能标识，与本软件设计说明第5章的'
                '操作处理设计行双向对照；训练主链已有的关键处理设计行同时列出。'
                '正向表逐条列需求，逆向表按设计项汇总需求标识。设计对应关系用于核对方案覆盖范围，'
                '不等于实现和验收已经完成。',
            )
            break


def digest(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


def document_sections(document):
    """Fingerprint text and table contents under each heading in the joint draft."""
    outline = {}
    occurrences = Counter()
    key = '文档前置内容'
    sections = {key: []}
    for element in document.element.body:
        if element.tag == qn('w:p'):
            paragraph = Paragraph(element, document._body)
            style = paragraph.style.name if paragraph.style is not None else ''
            if style.startswith('Heading ') and style.removeprefix('Heading ').isdigit():
                level = int(style.removeprefix('Heading '))
                outline = {depth: title for depth, title in outline.items() if depth < level}
                outline[level] = paragraph.text.strip()
                path = ' / '.join(title for _, title in sorted(outline.items()))
                occurrences[path] += 1
                key = f'{path} [{occurrences[path]}]'
                sections[key] = []
            else:
                sections[key].append(paragraph.text.strip())
        elif element.tag == qn('w:tbl'):
            sections[key].append(' | '.join(node.text or '' for node in element.iter(qn('w:t'))))
    return {key: digest('\n'.join(lines)) for key, lines in sections.items()}


def source_label(path):
    try:
        return str(path.resolve().relative_to(ROOT.parent.resolve()))
    except ValueError:
        return str(path.resolve())


def merge_register(source, target, document, sections, anchors, decisions, decisions_path,
                   model_signatures, structured_source, structured_path, read_model_signatures, interface_records,
                   requirement_refs, traceability_refs, corrections, corrections_path):
    return {
        'schemaVersion': 8,
        'source': source_label(source),
        'decisionSource': source_label(decisions_path),
        'structuredSource': source_label(structured_path),
        'correctionsSource': source_label(corrections_path),
        'target': source_label(target),
        'targetSha256': hashlib.sha256(target.read_bytes()).hexdigest(),
        'sourceSections': {key: digest(value) for key, value in sections.items()},
        'decisionItems': {
            item['id']: digest(json.dumps(item, ensure_ascii=False, sort_keys=True))
            for item in (*decisions['items'], *decisions.get('behaviorItems', []))
        },
        'decisionModelRefs': model_signatures,
        'structuredItems': structured_section_items(structured_source),
        'interfaceSourceRefs': interface_source_signatures(structured_source),
        'implementationSourceRefs': implementation_source_signatures(structured_source),
        'interfaceInventoryItems': {
            item['id']: digest(json.dumps(item, ensure_ascii=False, sort_keys=True))
            for item in interface_records
        },
        'readModelItems': read_model_signatures,
        'requirementRefs': requirement_refs,
        'traceabilityRefs': traceability_refs,
        'correctionItems': correction_signatures(corrections),
        'targetSections': document_sections(document),
        'anchors': {key: {'text': paragraph.text.strip(), 'style': paragraph.style.name}
                    for key, paragraph in anchors.items()},
    }


def changed_entries(previous, current):
    return [key for key in current if previous.get(key) != current[key]] + [
        key for key in previous if key not in current
    ]


def register_changes(previous, current):
    changes = []
    for field, label in (
        ('decisionItems', '输入/输出设计项'),
        ('decisionModelRefs', '关联建模定义'),
        ('structuredItems', '其他设计章节条目'),
        ('interfaceSourceRefs', '接口实现来源'),
        ('implementationSourceRefs', '详细设计实现来源'),
        ('interfaceInventoryItems', 'HTTP 接口清单'),
        ('readModelItems', '查询视图建模定义'),
        ('requirementRefs', '功能需求追踪'),
        ('traceabilityRefs', '第6章需求可追踪性'),
        ('correctionItems', '总稿受控勘误'),
    ):
        changed = changed_entries(previous.get(field, {}), current.get(field, {}))
        if changed:
            changes.append(f'{label}变更 {len(changed)} 处：{"；".join(changed[:12])}{"；…" if len(changed) > 12 else ""}')
    for field, label in (('sourceSections', '联邦学习正文'), ('targetSections', '总软设章节')):
        changed = changed_entries(previous.get(field, {}), current[field])
        if changed:
            sample = '；'.join(changed[:12])
            changes.append(f'{label}变更 {len(changed)} 处：{sample}{"；…" if len(changed) > 12 else ""}')
    if previous.get('target') != current['target']:
        changes.append(f'总软设底稿路径变更：{previous.get("target")} -> {current["target"]}')
    if previous.get('source') != current['source']:
        changes.append(f'联邦学习源稿路径变更：{previous.get("source")} -> {current["source"]}')
    if previous.get('decisionSource') != current.get('decisionSource'):
        changes.append(f'输入/输出设计源路径变更：{previous.get("decisionSource")} -> {current["decisionSource"]}')
    if previous.get('structuredSource') != current.get('structuredSource'):
        changes.append(f'其他设计章节源路径变更：{previous.get("structuredSource")} -> {current["structuredSource"]}')
    if previous.get('correctionsSource') != current.get('correctionsSource'):
        changes.append(f'总稿勘误源路径变更：{previous.get("correctionsSource")} -> {current["correctionsSource"]}')
    if previous.get('targetSha256') != current['targetSha256'] and not any(
        item.startswith('总软设章节变更') for item in changes
    ):
        changes.append('总软设文件已变更，但章节文字未变；请检查格式、图表或域。')
    return changes


def insert_section(document, anchor, markdown, figures=None):
    fragment = Document()
    write_markdown(fragment, markdown, '')
    style_ids = {
        style.style_id: document.styles[style.name].style_id
        for style in fragment.styles
        if style.name in document.styles
    }
    last = anchor._p
    for element in fragment.element.body:
        if element.tag not in {qn('w:p'), qn('w:tbl')}:
            continue
        if element.tag == qn('w:p'):
            marker = FIGURE_MARKER.fullmatch(''.join(node.text or '' for node in element.iter(qn('w:t'))))
            if marker:
                if not figures or marker.group(1) not in figures:
                    raise ValueError(f'Unknown software design figure: {marker.group(1)}')
                block, image_path = figures[marker.group(1)]
                image_paragraph = document.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.paragraph_format.keep_with_next = True
                image_paragraph.add_run().add_picture(str(image_path), width=Cm(block['widthCm']))
                last.addnext(image_paragraph._p)
                last = image_paragraph._p
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.add_run('图：' + block['caption']).font.size = Pt(9)
                last.addnext(caption._p)
                last = caption._p
                continue
        copied = deepcopy(element)
        for style_ref in copied.xpath('./w:pPr/w:pStyle | ./w:tblPr/w:tblStyle'):
            original_id = style_ref.get(qn('w:val'))
            if original_id not in style_ids:
                if style_ref.tag == qn('w:tblStyle'):
                    style_ref.getparent().remove(style_ref)
                    continue
                raise ValueError(f'Target document has no matching paragraph style for {original_id}')
            style_ref.set(qn('w:val'), style_ids[original_id])
        for fonts in copied.iter(qn('w:rFonts')):
            font = '黑体' if copied.tag == qn('w:p') and copied.xpath('./w:pPr/w:pStyle') else '宋体'
            for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn(f'w:{attribute}'), font)
        last.addnext(copied)
        last = copied


def insert_before_heading(document, heading, markdown):
    marker = document.add_paragraph()
    heading._p.addprevious(marker._p)
    try:
        insert_section(document, marker, markdown)
    finally:
        marker._p.getparent().remove(marker._p)


def format_inserted_tables(document, tables):
    """Use the joint design document's table style without changing existing tables."""
    reference = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == '序号')
    template_properties = reference._tbl.tblPr
    for table in tables:
        table._tbl.replace(table._tbl.tblPr, deepcopy(template_properties))
        table.autofit = False
        inventory_table = (
            len(table.columns) == 4
            and [cell.text.strip() for cell in table.rows[0].cells]
            == ['接口标识', '调用方', '方法与相对路径', '用途']
        )
        detailed_table = (
            len(table.columns) == 2
            and [cell.text.strip() for cell in table.rows[0].cells] == ['设计要素', '设计说明']
        )
        context_table = (
            len(table.columns) == 4
            and [cell.text.strip() for cell in table.rows[0].cells]
            == ['上下文软件单元', '输入与主要判断', '状态输出与异常', '读侧或部署边界']
        )
        operation_table = (
            len(table.columns) == 4
            and [cell.text.strip() for cell in table.rows[0].cells]
            == ['操作设计项', '输入与状态来源', '处理与外部依赖', '结果、异常及查询更新']
        )
        behavior_table = (
            len(table.columns) == 5
            and [cell.text.strip() for cell in table.rows[0].cells]
            == ['业务过程', '触发与输入', '状态与校验决策', '结果及异常', '查询展示']
        )
        if inventory_table:
            widths = [1500, 1500, 4500, 1500]
        elif detailed_table:
            widths = [2000, 7000]
        elif context_table:
            widths = [1700, 2550, 2550, 2200]
        elif operation_table:
            widths = [1300, 2400, 3100, 2200]
        elif behavior_table:
            widths = [1250, 1600, 2400, 2250, 1500]
        else:
            original_widths = [column.width.twips for column in table.columns]
            total_width = sum(original_widths)
            widths = [round(9000 * width / total_width) for width in original_widths[:-1]]
            widths.append(9000 - sum(widths))
        if inventory_table or detailed_table or context_table or operation_table or behavior_table:
            header = OxmlElement('w:tblHeader')
            header.set(qn('w:val'), 'true')
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for row in (table.rows[:1] if operation_table else table.rows):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for column, width in zip(table.columns, widths):
            column.width = Twips(width)
        for row_index, row in enumerate(table.rows):
            for cell, width in zip(row.cells, widths):
                cell.width = Twips(width)
                tc_pr = cell._tc.get_or_add_tcPr()
                for name in ('shd', 'tcBorders', 'tcMar'):
                    element = tc_pr.find(qn(f'w:{name}'))
                    if element is not None:
                        tc_pr.remove(element)
                for paragraph in cell.paragraphs:
                    paragraph.style = document.styles['RS-表-表头' if row_index == 0 else 'RS-表-内容']
                    paragraph.paragraph_format.space_after = None
                    paragraph.paragraph_format.line_spacing = None
                    paragraph.paragraph_format.alignment = None
                    for run in paragraph.runs:
                        run.font.size = None
                        run.bold = None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--target', type=Path, default=DEFAULT_TARGET)
    parser.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    parser.add_argument('--decisions', type=Path, default=DEFAULT_DECISIONS)
    parser.add_argument('--sections', type=Path, default=DEFAULT_SECTIONS)
    parser.add_argument('--corrections', type=Path, default=DEFAULT_CORRECTIONS)
    parser.add_argument('--out', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--sync', action='store_true', help='Accept the reviewed source and joint draft as a new merge baseline')
    args = parser.parse_args()
    if args.check and args.sync:
        parser.error('--check and --sync cannot be used together')
    if args.out.resolve() == args.target.resolve():
        raise ValueError('The output must be a copy; refusing to overwrite the joint source document.')

    sections = parse_sections(args.source.read_text(encoding='utf-8'))
    if sections['decisions'].count(DECISIONS_PLACEHOLDER) != 1:
        raise ValueError('Design merge source must contain the input/output decisions placeholder exactly once.')
    decisions = read_decisions(args.decisions)
    model = json.loads(MODEL.read_text(encoding='utf-8'))
    model_signatures = decision_model_signatures(decisions, model)
    sections['decisions'] = sections['decisions'].replace(DECISIONS_PLACEHOLDER, render_decisions(decisions))
    structured_source = read_structured_sections(args.sections)
    corrections = read_corrections(args.corrections)
    figures = figure_assets(structured_source)
    requirement_refs = requirement_signatures(structured_source)
    operation_groups = load_operation_design()
    traces = traceability_rows(structured_source, operation_groups=operation_groups)
    rendered_sections = render_structured_sections(structured_source, operation_groups)
    for key, rendered in rendered_sections.items():
        if sections[key].strip() != SECTION_PLACEHOLDER:
            raise ValueError(f'Software design section must use the structured placeholder: {key}')
        sections[key] = rendered
    read_models, read_model_signatures = read_model_inventory()
    sections['db-logical'] = sections['db-logical'].replace('{{READ_MODEL_INVENTORY}}', read_models)
    translations = model.get('translations', {}).get('zh-CN', {})
    interface_records = discover_routes(ROOT, translations)
    if sections['interfaces'].count('{{INTERFACE_INVENTORY}}') != 1:
        raise ValueError('Interface design must contain one complete HTTP inventory placeholder.')
    sections['interfaces'] = sections['interfaces'].replace(
        '{{INTERFACE_INVENTORY}}', render_inventory(interface_records)
    )
    document = Document(args.target)
    if any(paragraph.text.startswith('联邦学习子系统按平台、参与方运行时代理') for paragraph in document.paragraphs):
        raise ValueError('The target already contains this merge content; use the original joint document.')
    anchors = find_anchors(document)
    targets = correction_targets(document, corrections)
    current = merge_register(
        args.source, args.target, document, sections, anchors, decisions, args.decisions,
        model_signatures, structured_source, args.sections, read_model_signatures, interface_records,
        requirement_refs, traceability_signatures(traces), corrections, args.corrections,
    )
    previous = json.loads(REGISTER.read_text(encoding='utf-8')) if REGISTER.exists() else {}
    changes = register_changes(previous, current)
    for change in changes:
        print(change)
    if args.sync:
        REGISTER.write_text(json.dumps(current, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f'Accepted merge baseline: {REGISTER}')
        return
    if changes:
        raise ValueError('Merge baseline has changed; review the locations above, then run --sync to accept them.')
    if args.check:
        print(f'Checked {len(anchors)} design insertion points in {args.target}')
        return
    apply_corrections(corrections, targets)
    existing_tables = {table._tbl for table in document.tables}
    for key in ANCHORS:
        if key == 'performance':
            insert_section(document, anchors[key], OTHER_MODULES_PERFORMANCE)
            insert_before_heading(document, anchors['interfaces'], sections[key])
        else:
            insert_section(document, anchors[key], sections[key], figures)
    update_traceability_chapter(document, traces)
    added_tables = [table for table in document.tables if table._tbl not in existing_tables]
    format_inserted_tables(document, added_tables)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.out)
    print(f'Merged design copy: {args.out}')


if __name__ == '__main__':
    main()
