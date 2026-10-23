"""Regression checks for the project database-design merge."""

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from database_design_inventory import load_catalog, render_logical, render_physical, traceability_rows

spec = importlib.util.spec_from_file_location('project_document_merge', ROOT / 'scripts/merge-438c-project-documents.py')
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)


class DatabaseDesignMergeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_catalog()
        cls.register = json.loads((ROOT / 'docs/438c/merge/3.3.1-function-register.json').read_text(encoding='utf-8'))
        cls.target = merge.DATABASE_TARGET

    def test_catalog_matches_model_entities_and_requirements(self):
        self.assertEqual(56, len(self.catalog))
        self.assertEqual(56, len({item['table'] for item in self.catalog}))
        self.assertTrue(all(item['properties'] and item['fields'] for item in self.catalog))
        self.assertEqual(56, len(traceability_rows(self.catalog, self.register)))
        logical_inventory, logical_details = render_logical(self.catalog, self.register)
        physical_inventory, physical_details = render_physical(self.catalog)
        self.assertIn('LBLS-GN-LBXX-', logical_inventory)
        self.assertIn('FL-DB-L042', logical_details)
        self.assertIn('FL-DB-L056', logical_details)
        self.assertIn('FL-DB-P042', physical_details)
        self.assertIn('FL-DB-P056', physical_details)
        self.assertIn('organization_directory', physical_inventory)
        reversed_inventory, _ = render_logical(list(reversed(self.catalog)), self.register)
        self.assertIn('FL-DB-L001', reversed_inventory)
        self.assertIn('FL-DB-L042', reversed_inventory)
        self.assertIn('FL-DB-L056', reversed_inventory)

    def test_source_contains_controlled_sections(self):
        sections = merge.sections('database-design-federation-learning.md', {
            'conceptual', 'logical', 'physical', 'software-units',
        })
        self.assertEqual(4, len(sections))
        merge.merge_database(self.target, Path('/unused'), check=True)

    def test_merged_copy_preserves_original_and_adds_database_design(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'database-design.docx'
            merge.merge_database(self.target, output, check=False)
            document = Document(output)
        text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn('联邦学习模块数据库级设计决策', text)
        self.assertNotIn('联邦学习模块数据库需求追溯', text)
        self.assertNotIn('联邦学习模块正向追溯', text)
        self.assertNotIn('联邦学习模块逆向追溯', text)
        self.assertIn('FL-DB-L042', text)
        self.assertIn('FL-DB-L056', text)
        self.assertIn('FL-DB-P042', text)
        self.assertIn('FL-DB-P056', text)
        base = Document(self.target)
        def decision_paragraphs(doc):
            paragraphs = doc.paragraphs
            index = next(index for index, paragraph in enumerate(paragraphs)
                         if paragraph.text == '数据库级设计决策') + 1
            result = []
            while index < len(paragraphs) and paragraphs[index].style.name != 'Heading 1':
                result.append(paragraphs[index].text)
                index += 1
            return result
        self.assertEqual(decision_paragraphs(base), decision_paragraphs(document))
        self.assertNotIn('{{READ_MODEL_', text)
        headings = [paragraph for paragraph in document.paragraphs
                    if paragraph.style.name.startswith('Heading ')]
        for index, heading in enumerate(headings[:-1]):
            if heading.style.name == 'Heading 3' and heading.text == '联邦学习模块':
                self.assertEqual('Heading 4', headings[index + 1].style.name)
        for title in ('概念关系图', '核心数据对象', '训练链路概念关系图',
                      '查询视图目录', '查询视图数据元素', '查询表与主键映射',
                      '查询表属性设计'):
            self.assertEqual('Heading 4', next(heading.style.name for heading in headings
                                              if heading.text == title))
        self.assertEqual('Heading 2', next(heading.style.name for heading in headings
                                           if heading.text == '联邦学习模块数据库访问软件单元'))
        self.assertEqual('Heading 2', next(heading.style.name for heading in headings
                                           if heading.text == '跨域健康管理联邦孪生平台的软件单元设计'))
        self.assertNotIn('跨域健康管理联邦孪生平台的软件单元设计见表 117。', text)
        software_heading = next(heading for heading in headings
                                if heading.text == '联邦学习模块数据库访问软件单元')
        joint_software_table = next(table for table in document.tables if len(table.columns) == 6
                                    and table.rows[0].cells[0].text == '软件单元名称')
        federation_software_table = next(table for table in document.tables if len(table.columns) == 4
                                         and table.rows[0].cells[0].text == '软件单元标识')
        body_elements = list(document.element.body)
        self.assertLess(body_elements.index(joint_software_table._tbl),
                        body_elements.index(software_heading._p))
        self.assertLess(body_elements.index(software_heading._p),
                        body_elements.index(federation_software_table._tbl))
        self.assertGreaterEqual(len(document.inline_shapes), 2)
        trace_tables = [table for table in document.tables if len(table.columns) == 4
                        and len(table.rows) == 82 and table.rows[0].cells[0].text == '序号']
        self.assertEqual(2, len(trace_tables))
        self.assertFalse(any(paragraph.style.name == 'RS-表-题注' and not paragraph.text.strip()
                             for paragraph in document.paragraphs))
        base_trace_tables = [table for table in base.tables if len(table.columns) == 4
                             and len(table.rows) == 36 and table.rows[0].cells[0].text == '序号']
        self.assertEqual(2, len(base_trace_tables))
        for table, base_table in zip(trace_tables, base_trace_tables):
            self.assertFalse(table.autofit)
            self.assertEqual([650, 3300, 4500, 616], [column.width.twips for column in table.columns])
            self.assertEqual([str(index) for index in range(1, 82)],
                             [row.cells[0].text for row in table.rows[1:]])
            self.assertEqual([[cell.text for cell in row.cells[1:]] for row in base_table.rows[11:]],
                             [[cell.text for cell in row.cells[1:]] for row in table.rows[57:]])
        for index, (requirement_id, title, logical_id, physical_id, _table_name) in enumerate(
                traceability_rows(self.catalog, self.register), 1):
            requirement = f'{title}/{requirement_id}'
            design = f'{logical_id}；{physical_id}'
            self.assertEqual(requirement, trace_tables[0].rows[index].cells[1].text)
            self.assertEqual(design, trace_tables[0].rows[index].cells[2].text)
            self.assertEqual(design, trace_tables[1].rows[index].cells[1].text)
            self.assertEqual(requirement, trace_tables[1].rows[index].cells[2].text)
            self.assertEqual(base_trace_tables[0].rows[1].cells[1].paragraphs[0].alignment,
                             trace_tables[0].rows[index].cells[1].paragraphs[0].alignment)
            self.assertEqual(base_trace_tables[0].rows[1].cells[1].paragraphs[0].runs[0].font.size,
                             trace_tables[0].rows[index].cells[2].paragraphs[0].runs[0].font.size)
        inventory = next(table for table in document.tables if table.rows
                         and table.rows[0].cells[0].text == '设计标识')
        self.assertFalse(inventory.autofit)
        self.assertEqual('RS-表-表头', inventory.rows[0].cells[0].paragraphs[0].style.name)
        self.assertEqual('RS-表-内容', inventory.rows[1].cells[0].paragraphs[0].style.name)
        self.assertIsNotNone(inventory.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')))


if __name__ == '__main__':
    unittest.main()
